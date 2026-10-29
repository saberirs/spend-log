#!/usr/bin/env python3
"""
Spenderrr – automated test suite
================================

Covers Android, iPhone, Firefox-style browsers and desktop:
release/version checks, naming, security rule, privacy page, layout,
folder backups, monthly folders, the 6 PM schedule, iPhone Files / download
fallbacks, Google Drive backup and restore, restore from file and folder,
the safety checks (empty device, much smaller data), updates and offline use.

Google is never contacted: sign-in and Drive are simulated.

Setup (once):
    pip install playwright openpyxl
    python -m playwright install chromium

Run:
    python tests/run_tests.py                     # tests folder inside the app folder
    python run_tests.py --app ~/Documents/spenderrr   # tests kept elsewhere
    python tests/run_tests.py drive               # only tests whose name contains "drive"
    python tests/run_tests.py --headed            # watch the browser
    python tests/run_tests.py --verbose           # full error details
"""
import base64
import io
import json
import os
import zipfile
import re
import shutil
import sys
import tempfile
import threading
import time
import traceback
from datetime import datetime
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import unquote, urlparse

from playwright.sync_api import sync_playwright
import openpyxl

def _app_dir():
    if '--app' in sys.argv:
        return os.path.abspath(sys.argv[sys.argv.index('--app') + 1])
    return os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))


APP_DIR = _app_dir()
HEADED = '--headed' in sys.argv
_skip = {sys.argv[sys.argv.index('--app') + 1]} if '--app' in sys.argv else set()
FILTERS = [a for a in sys.argv[1:] if not a.startswith('--') and a not in _skip]

DATA_KEY = 'spendlog.v1'
META_KEY = '__spendlog_backup_meta'
EXPECTED_NAME = 'Spenderrr'
OLD_NAMES = ['Daily Spend Log', 'Spearmint Trackr', 'Spend Log']

SIZES = {
    'android': {'width': 412, 'height': 860},
    'iphone': {'width': 390, 'height': 844},
    'desktop': {'width': 1366, 'height': 800},
}


# --------------------------------------------------------------------------
# Local web server that behaves like Netlify (applies the _headers file)
# --------------------------------------------------------------------------
def parse_headers_file(path):
    rules, current = [], None
    if not os.path.exists(path):
        return rules
    for raw in open(path, encoding='utf-8'):
        line = raw.rstrip('\n')
        if not line.strip() or line.lstrip().startswith('#'):
            continue
        if not line.startswith((' ', '\t')):
            current = {'pattern': line.strip(), 'headers': {}}
            rules.append(current)
        elif current is not None and ':' in line:
            k, v = line.strip().split(':', 1)
            current['headers'][k.strip()] = v.strip()
    return rules


def pattern_matches(pattern, path):
    if pattern.endswith('*'):
        return path.startswith(pattern[:-1])
    return path == pattern


STUB_XLSX = b"window.XLSX={utils:{book_new:()=>({SheetNames:[],Sheets:{}}),aoa_to_sheet:()=>({}),json_to_sheet:()=>({}),book_append_sheet:()=>{}},writeFile:()=>{}};"
PNG_1PX = bytes.fromhex('89504e470d0a1a0a0000000d4948445200000001000000010806000000'
                        '1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082')


class AppServer:
    def __init__(self, root):
        self.root = root
        server = self

        class Handler(SimpleHTTPRequestHandler):
            def translate_path(self, path):
                p = unquote(urlparse(path).path)
                if p.endswith('/'):
                    p += 'index.html'
                return os.path.join(server.root, p.lstrip('/'))

            def end_headers(self):
                p = urlparse(self.path).path
                for rule in parse_headers_file(os.path.join(server.root, '_headers')):
                    if pattern_matches(rule['pattern'], p):
                        for k, v in rule['headers'].items():
                            self.send_header(k, v)
                super().end_headers()

            def do_GET(self):
                p = urlparse(self.path).path
                local = self.translate_path(self.path)
                if not os.path.exists(local):   # files not in the test folder (library, icons)
                    if p.endswith('xlsx.full.min.js'):
                        return self._send(STUB_XLSX, 'text/javascript')
                    if p.endswith('.png'):
                        return self._send(PNG_1PX, 'image/png')
                return super().do_GET()

            def _send(self, body, ctype):
                self.send_response(200)
                self.send_header('Content-Type', ctype)
                self.send_header('Content-Length', str(len(body)))
                self.end_headers()
                self.wfile.write(body)

            def log_message(self, *a):
                pass

        Handler.extensions_map = dict(SimpleHTTPRequestHandler.extensions_map,
                                      **{'.webmanifest': 'application/manifest+json', '.js': 'text/javascript'})
        self.httpd = ThreadingHTTPServer(('127.0.0.1', 0), Handler)
        self.url = 'http://127.0.0.1:%d' % self.httpd.server_address[1]
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def stop(self):
        self.httpd.shutdown()


# --------------------------------------------------------------------------
# Simulated Google sign-in and Google Drive
# --------------------------------------------------------------------------
FAKE_GIS = """
window.google = { accounts: { oauth2: {
  initTokenClient: function (c) { return { requestAccessToken: function () {
    window.__gisRequests = (window.__gisRequests || 0) + 1;
    setTimeout(function () { c.callback({ access_token: 'test-token', expires_in: 3600 }); }, 10);
  } }; },
  revoke: function (t, cb) { cb && cb(); }
} } };
"""
FOLDER = 'application/vnd.google-apps.folder'


TEST_ACCOUNT = {'displayName': 'Saberi R', 'emailAddress': 'saberi.test@gmail.com'}


class FakeDrive:
    """A small stand-in for Google Drive: folders, files, newest-first sorting, trash, account lookup.
    Several Device objects can share one FakeDrive = several phones using the same Google account."""

    def __init__(self, account=None):
        self.account = dict(account or TEST_ACCOUNT)   # set to None to make Google refuse the account lookup
        self.files = {}
        self.calls = []
        self._n = 0
        self._clock = 0

    def _id(self):
        self._n += 1
        return 'f%d' % self._n

    def _tick(self):
        self._clock += 1
        return '2026-01-01T00:00:%06dZ' % self._clock   # strictly increasing timestamps

    def add(self, name, parent, mime, body=None):
        i, t = self._id(), self._tick()
        self.files[i] = {'name': name, 'parents': [parent], 'mime': mime, 'body': body,
                         'trashed': False, 'createdTime': t, 'modifiedTime': t}
        return i

    def live(self):
        return {k: f for k, f in self.files.items() if not f['trashed']}

    def tree(self):
        """{'spend-expense-backup/October_2026/spend-log-backup.json': body} — files not in the trash."""
        def path(i):
            f = self.files[i]
            par = f['parents'][0]
            return f['name'] if par == 'root' else path(par) + '/' + f['name']
        return {path(i): f['body'] for i, f in self.live().items() if f['mime'] != FOLDER}

    def copies(self, name=None):
        """Every live file with this name (duplicates included), newest first: [(id, body)]."""
        hits = [(k, f) for k, f in self.live().items() if f['mime'] != FOLDER and (name is None or f['name'] == name)]
        hits.sort(key=lambda kf: kf[1]['modifiedTime'], reverse=True)
        return [(k, f['body']) for k, f in hits]

    def handle(self, route):
        req, url, m = route.request, route.request.url, route.request.method
        self.calls.append(m + ' ' + url.split('googleapis.com')[-1][:90])
        ok = lambda body: route.fulfill(status=200, content_type='application/json', body=json.dumps(body))
        if '/drive/v3/about' in url and m == 'GET':
            if self.account is None:
                return route.fulfill(status=403, content_type='application/json', body='{"error":"forbidden"}')
            return ok({'user': self.account})
        if '/drive/v3/files?' in url and m == 'GET':
            q = unquote(re.search(r'q=([^&]+)', url).group(1))
            nm = re.search(r"name='([^']+)'", q)
            par = re.search(r"'([^']+)' in parents", q).group(1)
            want_folder = "mimeType='" in q
            order = unquote(re.search(r'orderBy=([^&]+)', url).group(1)) if 'orderBy=' in url else ''
            hits = [(k, f) for k, f in self.live().items()
                    if (nm is None or f['name'] == nm.group(1)) and par in f['parents'] and (f['mime'] == FOLDER) == want_folder]
            if order.startswith('modifiedTime'):
                hits.sort(key=lambda kf: kf[1]['modifiedTime'], reverse='desc' in order)
            elif order.startswith('createdTime'):
                hits.sort(key=lambda kf: kf[1]['createdTime'], reverse='desc' in order)
            elif order.startswith('name'):
                hits.sort(key=lambda kf: kf[1]['name'])
            return ok({'files': [{'id': k, 'name': f['name'], 'modifiedTime': f['modifiedTime'],
                                  'createdTime': f['createdTime']} for k, f in hits]})
        if '/drive/v3/files?fields=id' in url and m == 'POST':
            d = json.loads(req.post_data)
            return ok({'id': self.add(d['name'], d['parents'][0], d['mimeType'])})
        if 'uploadType=multipart' in url and m == 'POST':
            parts = req.post_data.split('\r\n\r\n')
            meta = json.loads(parts[1].split('\r\n--')[0])
            body = parts[2].rsplit('\r\n--', 1)[0]
            return ok({'id': self.add(meta['name'], meta['parents'][0], meta['mimeType'], body)})
        if '/upload/drive/v3/files/' in url and m == 'PATCH':
            i = url.split('/files/')[1].split('?')[0]
            binary = 'spreadsheetml' in (req.headers.get('content-type') or '')
            self.files[i]['body'] = req.post_data_buffer if binary else req.post_data
            self.files[i]['modifiedTime'] = self._tick()
            return ok({})
        if '/drive/v3/files/' in url and m == 'PATCH' and '/upload/' not in url:
            i = url.split('/files/')[1].split('?')[0]
            d = json.loads(req.post_data or '{}')
            if d.get('trashed'):
                self.files[i]['trashed'] = True
            return ok({'id': i})
        if 'alt=media' in url:
            i = url.split('/files/')[1].split('?')[0]
            return route.fulfill(status=200, content_type='application/json', body=self.files[i]['body'])
        return route.fulfill(status=404, body='not mocked')


# --------------------------------------------------------------------------
# Device profiles (injected before the app loads)
# --------------------------------------------------------------------------
COMMON_INIT = """
window.__csp = [];
document.addEventListener('securitypolicyviolation', function (e) {
  window.__csp.push(e.violatedDirective + ' ' + e.blockedURI);
});
"""
# Android Chrome / desktop: a "Documents" folder the app can write to (a private test folder).
# Test browsers crash when a private test folder is saved in IndexedDB, so the folder the app
# remembers is stored as a marker and handed back on read, the way a phone remembers a real folder.
FOLDER_INIT = """
(function () {
  var MARK = 'test-folder:Documents';
  window.__docs = null;
  window.__docsReady = navigator.storage.getDirectory()
    .then(function (r) { return r.getDirectoryHandle('Documents', { create: true }); })
    .then(function (h) { window.__docs = h; return h; });
  Object.defineProperty(window, 'showDirectoryPicker', { configurable: true, writable: true, value: async function () {
    window.__pickerCalls = Number(sessionStorage.getItem('__pickerCalls') || 0) + 1;
    sessionStorage.setItem('__pickerCalls', window.__pickerCalls);
    return window.__docsReady;
  } });
  window.__pickerCalls = Number(sessionStorage.getItem('__pickerCalls') || 0);
  var put = IDBObjectStore.prototype.put, get = IDBObjectStore.prototype.get;
  var resultOf = Object.getOwnPropertyDescriptor(IDBRequest.prototype, 'result').get;
  IDBObjectStore.prototype.put = function (value, key) {
    if (typeof FileSystemDirectoryHandle !== 'undefined' && value instanceof FileSystemDirectoryHandle) value = MARK;
    return put.call(this, value, key);
  };
  IDBObjectStore.prototype.get = function (key) {
    var req = get.call(this, key);
    Object.defineProperty(req, 'result', { get: function () {
      var v = resultOf.call(req); return v === MARK ? window.__docs : v;
    } });
    return req;
  };
})();
"""
# iPhone Safari: no folder access, share sheet with "Save to Files".
IPHONE_INIT = """
Object.defineProperty(window, 'showDirectoryPicker', { configurable: true, value: undefined });
window.__shared = [];
Object.defineProperty(navigator, 'canShare', { configurable: true, value: function (d) { return !!(d && d.files); } });
Object.defineProperty(navigator, 'share', { configurable: true, value: async function (d) {
  for (var i = 0; i < d.files.length; i++) window.__shared.push({ name: d.files[i].name, text: await d.files[i].text() });
} });
"""
# Firefox-style: no folder access and no share sheet, so the backup downloads.
FIREFOX_INIT = """
Object.defineProperty(window, 'showDirectoryPicker', { configurable: true, value: undefined });
Object.defineProperty(navigator, 'canShare', { configurable: true, value: undefined });
Object.defineProperty(navigator, 'share', { configurable: true, value: undefined });
"""
READ_FOLDER_JS = """async () => {
  const out = {};
  async function walk(dir, prefix) {
    for await (const [name, h] of dir.entries()) {
      const p = prefix ? prefix + '/' + name : name;
      if (h.kind === 'directory') { await walk(h, p); continue; }
      const f = await h.getFile();
      if (name.endsWith('.xlsx')) {
        const a = new Uint8Array(await f.arrayBuffer()); let s = '';
        for (let i = 0; i < a.length; i += 8192) s += String.fromCharCode.apply(null, a.subarray(i, i + 8192));
        out[p] = 'base64:' + btoa(s);
      } else out[p] = await f.text();
    }
  }
  const root = await navigator.storage.getDirectory();
  try { await walk(await root.getDirectoryHandle('Documents'), ''); } catch (e) {}
  return out;
}"""


class Device:
    """One browser profile = one phone/computer, with its own storage."""

    def __init__(self, browser, server, kind='android', size=None, clock=None, drive=None, sw=False):
        self.server, self.kind = server, kind
        self.ctx = browser.new_context(viewport=SIZES[size or {'desktop': 'desktop', 'iphone': 'iphone'}.get(kind, 'android')],
                                       service_workers='allow' if sw else 'block', accept_downloads=True)
        self.ctx.add_init_script(COMMON_INIT)
        self.ctx.add_init_script({'android': FOLDER_INIT, 'desktop': FOLDER_INIT,
                                  'iphone': IPHONE_INIT, 'firefox': FIREFOX_INIT}[kind])
        self.ctx.route('https://fonts.googleapis.com/**', lambda r: r.fulfill(status=200, content_type='text/css', body=''))
        self.ctx.route('https://fonts.gstatic.com/**', lambda r: r.fulfill(status=404, body=''))
        self.ctx.route('https://accounts.google.com/gsi/client',
                       lambda r: r.fulfill(status=200, content_type='text/javascript', body=FAKE_GIS))
        self.drive = drive or FakeDrive()
        self.ctx.route('https://www.googleapis.com/**', self.drive.handle)
        self.page = self.ctx.new_page()
        self.dialogs, self.answers = [], []
        self.page.on('dialog', self._dialog)
        self.errors = []
        self.page.on('pageerror', lambda e: self.errors.append(str(e)))
        if clock:
            self.page.clock.install(time=clock)

    def _dialog(self, d):
        self.dialogs.append(d.message)
        ans = self.answers.pop(0) if self.answers else True
        d.accept() if ans else d.dismiss()

    # navigation / helpers
    def open(self, path='/'):
        self.page.goto(self.server.url + path)
        self.page.wait_for_selector('#slb-fab')
        self.page.wait_for_timeout(300)
        return self

    def reload(self):
        try:
            self.page.reload()
        except Exception:
            pass
        self.page.wait_for_selector('#slb-fab')
        self.page.wait_for_timeout(300)

    def ev(self, js, arg=None):
        return self.page.evaluate(js, arg) if arg is not None else self.page.evaluate(js)

    def records(self):
        return self.ev('SpendLogBackup.records()')

    def add_expense(self, amount, category='Lunch'):
        pg = self.page
        pg.click('#fab')
        pg.wait_for_selector('#fAmt')
        pg.fill('#fAmt', str(amount))
        pg.click('#sheetRoot button:has-text("%s")' % category)
        pg.click('#sheetRoot button[type=submit]')
        pg.wait_for_timeout(300)

    def seed(self, n):
        """Put n entries on this device (as the app stores them) and reload."""
        self.ev("""(n) => {
          const e = [];
          for (let i = 0; i < n; i++) e.push({ id: 'seed' + i, amount: 10 + i, cat: 'lunch', sub: null, note: '',
            ts: Date.now() - i * 3600000, source: 'manual', createdAt: Date.now(), split: null });
          localStorage.setItem('spendlog.v1', JSON.stringify({ e: e, m: [], p: [], s: [] }));
        }""", n)
        self.reload()

    def wipe_app_data(self):
        """What happens on reinstall / new home-screen icon / cleared site data. Files on the phone stay."""
        self.ev("""async () => { localStorage.clear(); sessionStorage.clear();
          await new Promise(r => { const q = indexedDB.deleteDatabase('spendlog-backup'); q.onsuccess = q.onerror = q.onblocked = r; }); }""")
        self.reload()

    def folder(self):
        return self.ev(READ_FOLDER_JS)

    def backup_now(self):
        self.ev('SpendLogBackup.backupNow()')
        self.page.wait_for_timeout(400)

    def run_auto(self):
        self.ev('SpendLogBackup.runAuto()')
        self.page.wait_for_timeout(400)

    def banner(self):
        b = self.page.query_selector('#slb-banner')
        return b.inner_text() if b and b.is_visible() else ''

    def meta(self):
        return json.loads(self.ev("localStorage.getItem('%s') || '{}'" % META_KEY))

    def close(self):
        assert not self.ev('window.__csp'), 'Security rule blocked: %s' % self.ev('window.__csp')
        assert not self.errors, 'Page errors: %s' % self.errors
        self.ctx.close()


def records_in(text):
    b = json.loads(text)
    o = json.loads(b['data'].get(DATA_KEY) or '{}')
    return sum(len(v) for v in o.values() if isinstance(v, list))


OCT = datetime(2026, 10, 26, 17, 0)        # Monday, before 6 PM
OCT_EVENING = datetime(2026, 10, 26, 19, 0)
MONTH_FILE = 'spend-expense-backup/October_2026/spend-log-backup.json'
MONTH_XLSX = 'spend-expense-backup/October_2026/spend-log-backup.xlsx'
NOV_FILE = 'spend-expense-backup/November_2026/spend-log-backup.json'
NOV_XLSX = 'spend-expense-backup/November_2026/spend-log-backup.xlsx'


# --------------------------------------------------------------------------
# Tests
# --------------------------------------------------------------------------
TESTS = []


def test(fn):
    TESTS.append(fn)
    return fn


def read(name):
    return open(os.path.join(APP_DIR, name), encoding='utf-8').read()


# ---- Release and configuration -------------------------------------------
@test
def release_version_is_the_same_everywhere(b, srv):
    html, sw = read('index.html'), read('sw.js')
    v = re.search(r'name="app-version" content="([^"]+)"', html).group(1)
    assert re.fullmatch(r'\d{4}\.\d{1,2}\.\d+', v), 'Version format should be YYYY.M.N, got ' + v
    assert 'backup.js?v=' + v in html, 'backup.js link must carry ?v=' + v
    assert 'whatsapp.js?v=' + v in html, 'whatsapp.js link must carry ?v=' + v
    assert "'/whatsapp.js'" in sw, 'sw.js must keep whatsapp.js for offline use'
    assert "spendlog-" + v in sw and 'Version ' + v in sw, 'sw.js must use version ' + v
    d = Device(b, srv).open()
    assert d.ev('SpendLogBackup.version') == v
    assert ('Version ' + v) in d.page.inner_text('footer.copy')
    d.close()


@test
def app_name_is_spenderrr_everywhere(b, srv):
    man = json.loads(read('manifest.webmanifest'))
    assert man['name'] == EXPECTED_NAME and man['short_name'] == EXPECTED_NAME
    html = read('index.html')
    assert '<title>%s</title>' % EXPECTED_NAME in html
    assert 'apple-mobile-web-app-title" content="%s"' % EXPECTED_NAME in html
    assert EXPECTED_NAME in read('privacy.html')
    d = Device(b, srv).open()
    d.seed(2)
    screens = []
    for view in ['recent', 'week', 'month', 'split', 'messages']:
        d.page.click('button[data-view=%s]' % view)
        d.page.wait_for_timeout(200)
        screens.append(d.ev('document.body.innerText'))
    d.page.click('#settingsBtn')
    d.page.wait_for_timeout(300)
    screens.append(d.ev('document.body.innerText'))
    text = '\n'.join(screens) + d.page.title()
    assert d.page.inner_text('h1') == EXPECTED_NAME
    for old in OLD_NAMES:
        assert old not in text, 'Old name "%s" still on screen' % old
    d.close()


@test
def security_rule_allows_google_and_nothing_breaks(b, srv):
    rules = parse_headers_file(os.path.join(APP_DIR, '_headers'))
    csp = next(r['headers']['Content-Security-Policy'] for r in rules if r['pattern'] == '/*')
    for needed in ['https://accounts.google.com/gsi/client', 'https://www.googleapis.com']:
        assert needed in csp, needed + ' missing from Content-Security-Policy'
    assert "frame-ancestors 'none'" in csp and "object-src 'none'" in csp, 'Existing protections were removed'
    assert "font-src 'self';" in csp and 'fonts.g' not in csp, 'Fonts come from the app itself; Google Fonts is not needed'
    d = Device(b, srv).open()
    loaded = d.ev("""() => new Promise(res => { const s = document.createElement('script');
        s.src = 'https://accounts.google.com/gsi/client'; s.onload = () => res(true); s.onerror = () => res(false);
        document.head.appendChild(s); })""")
    assert loaded, 'Google sign-in script is blocked'
    d.close()   # also asserts no security-rule violations happened


@test
def privacy_page_is_complete_and_linked(b, srv):
    d = Device(b, srv).open()
    d.page.click('footer.copy a[href="privacy.html"]')
    d.page.wait_for_load_state()
    t = d.ev('document.body.innerText')
    for must in ['drive.file', 'Limited Use', 'saberirajamanis@gmail.com', EXPECTED_NAME, 'account name and email']:
        assert must in t, 'Privacy page is missing: ' + must
    assert 'CONTACT_EMAIL' not in t
    d.ctx.close()


# ---- Layout on every device ------------------------------------------------
@test
def layout_same_on_android_iphone_desktop(b, srv):
    for size in SIZES:
        d = Device(b, srv, kind='desktop' if size == 'desktop' else 'android', size=size).open()
        r = d.ev("""() => { const g = id => document.getElementById(id).getBoundingClientRect();
            return { excel: g('exportBtn'), add: g('fab'), backup: g('slb-fab'),
                     scrollW: document.documentElement.scrollWidth, w: innerWidth, h: innerHeight }; }""")
        assert r['add']['top'] >= r['excel']['bottom'] - 1, size + ': Add should be below Excel'
        assert r['backup']['top'] >= r['excel']['bottom'] - 1, size + ': Backup should be below Excel'
        assert abs(r['backup']['right'] - r['excel']['right']) < 2, size + ': Backup should line up with Excel'
        assert r['add']['bottom'] < r['h'] - 120 and r['backup']['bottom'] < r['h'] - 120, \
            size + ': buttons must stay clear of the bottom bar'
        assert r['scrollW'] <= r['w'], size + ': page scrolls sideways'
        d.page.click('#slb-fab')
        panel = d.ev("(() => { const p = document.getElementById('slb-panel').getBoundingClientRect(); return [p.left, p.right]; })()")
        assert panel[0] >= 0 and panel[1] <= r['w'], size + ': backup panel goes off screen'
        d.close()


# ---- App basics -------------------------------------------------------------
@test
def entries_survive_reload(b, srv):
    d = Device(b, srv).open()
    d.add_expense(12.5, 'Lunch')
    d.add_expense(40, 'Grocery')
    d.reload()
    assert d.records() == 2
    assert '$52.50' in d.page.inner_text('main')
    d.close()


# ---- Folder backup (Android / desktop) -------------------------------------
@test
def folder_backup_creates_month_folder_and_file(b, srv):
    d = Device(b, srv, clock=OCT).open()
    d.seed(3)
    d.backup_now()
    files = d.folder()
    assert sorted(files) == [MONTH_FILE, MONTH_XLSX], files
    assert records_in(files[MONTH_FILE]) == 3
    assert d.ev('window.__pickerCalls') == 1
    assert d.meta()['last'] == '2026-10-26'
    d.close()


@test
def folder_backup_replaces_same_file_each_day(b, srv):
    d = Device(b, srv, clock=OCT).open()
    d.seed(3)
    d.backup_now()
    d.add_expense(9, 'Snacks')
    d.page.clock.set_system_time(datetime(2026, 10, 27, 19, 0))
    d.run_auto()
    files = d.folder()
    assert sorted(files) == [MONTH_FILE, MONTH_XLSX], 'Should still be one backup and one Excel file, got %s' % list(files)
    assert records_in(files[MONTH_FILE]) == 4
    assert d.ev('window.__pickerCalls') == 1, 'Folder should only be chosen once'
    d.close()


@test
def new_month_gets_new_folder(b, srv):
    d = Device(b, srv, clock=datetime(2026, 10, 31, 19, 0)).open()
    d.seed(3)
    d.backup_now()
    d.add_expense(5, 'Snacks')
    d.page.clock.set_system_time(datetime(2026, 11, 1, 19, 0))
    d.run_auto()
    files = d.folder()
    assert set(files) == {MONTH_FILE, MONTH_XLSX, NOV_FILE, NOV_XLSX}, sorted(files)
    assert records_in(files[MONTH_FILE]) == 3, 'October file must not change'
    assert records_in(files['spend-expense-backup/November_2026/spend-log-backup.json']) == 4
    d.close()


@test
def schedule_runs_after_6pm_once_a_day(b, srv):
    d = Device(b, srv, clock=datetime(2026, 10, 26, 19, 0)).open()
    d.seed(2)
    d.backup_now()                                   # yesterday's backup done
    d.page.clock.set_system_time(datetime(2026, 10, 27, 17, 0))
    assert d.ev('SpendLogBackup.isDue()') is False, 'Should not run before 6 PM'
    d.page.clock.set_system_time(datetime(2026, 10, 27, 18, 1))
    assert d.ev('SpendLogBackup.isDue()') is True, 'Should run after 6 PM'
    d.add_expense(7, 'Dinner')
    d.page.clock.run_for(61000)                      # the app checks every minute on its own
    d.page.wait_for_timeout(500)
    assert d.meta()['last'] == '2026-10-27'
    assert records_in(d.folder()[MONTH_FILE]) == 3
    assert d.ev('SpendLogBackup.isDue()') is False, 'Should not run twice the same day'
    d.close()


@test
def missed_day_catches_up_next_open(b, srv):
    d = Device(b, srv, clock=datetime(2026, 10, 24, 19, 0)).open()
    d.seed(2)
    d.backup_now()
    d.page.clock.set_system_time(datetime(2026, 10, 26, 9, 0))   # skipped the 25th, now morning
    assert d.ev('SpendLogBackup.isDue()') is True, 'A missed day should back up on next open'
    d.close()


# ---- iPhone and Firefox ----------------------------------------------------
@test
def iphone_evening_banner_and_save_to_files(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    d.seed(3)
    d.ev('SpendLogBackup.check()')
    d.page.wait_for_timeout(300)
    assert 'save to Files' in d.banner(), 'Banner: ' + d.banner()
    d.page.click('#slb-banner')
    d.page.wait_for_timeout(400)
    shared = d.ev('window.__shared')
    assert [f['name'] for f in shared] == ['spend-log-backup.json', 'spend-log-backup.xlsx'], [f['name'] for f in shared]
    assert records_in(shared[0]['text']) == 3
    assert not d.banner()
    d.close()


@test
def firefox_downloads_backup(b, srv):
    d = Device(b, srv, kind='firefox', clock=OCT_EVENING).open()
    d.seed(2)
    with d.page.expect_download() as dl:
        d.ev('SpendLogBackup.backupNow()')
    assert dl.value.suggested_filename == 'spend-expense-backup_October_2026.json'
    assert records_in(open(dl.value.path(), encoding='utf-8').read()) == 2
    d.close()


# ---- Google Drive -----------------------------------------------------------
@test
def drive_backup_creates_tree_and_replaces_file(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    d.seed(3)
    d.page.click('#slb-fab')
    d.page.click('text=Back up to Google Drive')
    d.page.wait_for_timeout(800)
    assert d.drive.tree() and records_in(d.drive.tree()[MONTH_FILE]) == 3, d.drive.tree()
    d.add_expense(4, 'Snacks')
    d.backup_now()
    assert sorted(d.drive.tree()) == [MONTH_FILE, MONTH_XLSX], 'Should be one backup and one Excel file on Drive'
    assert records_in(d.drive.tree()[MONTH_FILE]) == 4
    assert any(c.startswith('PATCH') for c in d.drive.calls), 'Second backup should replace the file'
    d.page.clock.set_system_time(datetime(2026, 11, 2, 19, 0))
    d.backup_now()
    assert 'spend-expense-backup/November_2026/spend-log-backup.json' in d.drive.tree()
    d.close()


@test
def drive_needs_a_tap_each_evening(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    d.seed(2)
    d.page.click('#slb-fab')
    d.page.click('text=Back up to Google Drive')
    d.page.wait_for_timeout(600)
    d.ev("sessionStorage.clear()")                    # sign-in expires (new day, app reopened)
    d.page.clock.set_system_time(datetime(2026, 10, 27, 19, 0))
    d.reload()
    d.ev('SpendLogBackup.check()')
    d.page.wait_for_timeout(300)
    assert 'Google Drive' in d.banner(), 'Banner: ' + d.banner()
    d.page.click('#slb-banner')
    d.page.wait_for_timeout(600)
    assert d.meta()['last'] == '2026-10-27'
    d.close()


@test
def restore_from_drive_on_new_phone(b, srv):
    drive = FakeDrive()
    old = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()
    old.seed(6)
    old.page.click('#slb-fab')
    old.page.click('text=Back up to Google Drive')
    old.page.wait_for_timeout(600)
    old.close()
    new = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()   # new phone, same Google account
    assert new.records() == 0
    new.page.click('#slb-banner')                         # "No entries ... Tap to restore"
    new.page.click('#slb-modal button:has-text("Google Drive")')
    new.page.wait_for_selector('#slb-modal .slb-item')
    items = new.page.locator('#slb-modal .slb-item')
    assert items.count() == 1 and '6 records' in items.first.inner_text(), items.first.inner_text()
    items.first.click()
    new.page.wait_for_timeout(800)
    new.reload()
    assert new.records() == 6
    assert new.meta().get('drive') is True, 'Drive should be connected after restoring from it'
    new.close()


# ---- Restore ----------------------------------------------------------------
@test
def restore_from_file_with_counts(b, srv):
    d = Device(b, srv, kind='firefox', clock=OCT_EVENING).open()
    d.seed(5)
    with d.page.expect_download() as dl:
        d.ev('SpendLogBackup.backupNow()')
    path = dl.value.path()
    d.seed(1)
    d.answers = [False]                                   # cancel first time: nothing changes
    d.page.click('#slb-fab')
    d.page.click('#slb-panel button:has-text("Restore")')
    with d.page.expect_file_chooser() as fc:
        d.page.click('#slb-modal button:has-text("Choose a file")')
    fc.value.set_files(path)
    d.page.wait_for_timeout(400)
    assert d.records() == 1, 'Cancel must keep current data'
    assert '5 records' in d.dialogs[-1] and 'replaces the 1 record on this device' in d.dialogs[-1], d.dialogs[-1]
    d.ev('SpendLogBackup.restore()')
    with d.page.expect_file_chooser() as fc:
        d.page.click('#slb-modal button:has-text("Choose a file")')
    fc.value.set_files(path)
    d.page.wait_for_timeout(800)
    d.reload()
    assert d.records() == 5
    d.close()


@test
def restore_from_folder_after_month_change(b, srv):
    d = Device(b, srv, clock=datetime(2026, 10, 30, 19, 0)).open()
    d.seed(4)
    d.backup_now()
    d.page.clock.set_system_time(datetime(2026, 11, 1, 9, 0))   # new month, no November backup yet
    d.wipe_app_data()
    d.page.click('#slb-banner')
    d.page.click('#slb-modal button:has-text("Backup folder")')
    d.page.wait_for_selector('#slb-modal .slb-item')
    first = d.page.locator('#slb-modal .slb-item').first
    assert 'October 2026' in first.inner_text() and '4 records' in first.inner_text(), first.inner_text()
    first.click()
    d.page.wait_for_timeout(800)
    d.reload()
    assert d.records() == 4, 'Should restore the October backup'
    d.close()


def put_in_folder(d, path, text):
    """Write a file into the test Documents folder, e.g. 'Spend-Expense-Backup/OCTOBER_2026/x.json'."""
    d.ev("""async (a) => {
      let dir = await (await navigator.storage.getDirectory()).getDirectoryHandle('Documents', { create: true });
      const parts = a.path.split('/'), file = parts.pop();
      for (const p of parts) dir = await dir.getDirectoryHandle(p, { create: true });
      const w = await (await dir.getFileHandle(file, { create: true })).createWritable();
      await w.write(a.text); await w.close();
    }""", {'path': path, 'text': text})


def restore_list(d, where):
    d.ev('SpendLogBackup.restore()')
    d.page.click('#slb-modal button:has-text("%s")' % where)
    d.page.wait_for_selector('#slb-modal .slb-item')
    loc = d.page.locator('#slb-modal .slb-item')
    return [loc.nth(i).inner_text() for i in range(loc.count())]


@test
def folder_backup_names_ignore_letter_case(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    put_in_folder(d, 'Spend-Expense-Backup/OCTOBER_2026/Spend-Log-Backup.JSON', backup_body([entry('a1', 10), entry('a2', 20), entry('a3', 30)]))
    put_in_folder(d, 'Spend-Expense-Backup/october_2026/spend-log-backup.json', backup_body([entry('a4', 40)], '2026-10-21T12:00:00Z'))
    put_in_folder(d, 'Spend-Expense-Backup/september_2026/SPEND-LOG-BACKUP.json', backup_body([entry('s1', 5), entry('s2', 6)]))
    put_in_folder(d, 'Spend-Expense-Backup/Notes/spend-log-backup.json', backup_body([entry('n1', 1)]))   # not a month
    items = restore_list(d, 'Backup folder')
    assert [i.split('\n')[0] for i in items] == ['October 2026', 'September 2026'], items
    assert '4 records' in items[0] and '2 records' in items[1], 'Both October folders count as one month: %s' % items
    d.page.locator('#slb-modal .slb-item').first.click()
    d.page.wait_for_timeout(800)
    d.reload()
    assert amounts(d) == {'a1': 10, 'a2': 20, 'a3': 30, 'a4': 40}, amounts(d)
    d.add_expense(7, 'Lunch')
    d.backup_now()
    files = d.folder()
    assert not any(k.startswith('spend-expense-backup/') or '/October_2026/' in k for k in files), \
        'Existing folders must be reused, not made again with the usual name: %s' % sorted(files)
    written = [k for k in files if k.lower().endswith('october_2026/spend-log-backup.json') and records_in(files[k]) == 5]
    assert len(written) == 1, 'The backup should go into an existing October file: %s' % sorted(files)
    assert written[0].rsplit('/', 1)[0] + '/spend-log-backup.xlsx' in files, 'Excel copy sits next to it'
    d.close()


@test
def drive_backup_names_ignore_letter_case(b, srv):
    drive = FakeDrive()
    root = drive.add('Spend-Expense-Backup', 'root', FOLDER)
    oct_a = drive.add('october_2026', root, FOLDER)
    oct_b = drive.add('OCTOBER_2026', root, FOLDER)
    sep = drive.add('September_2026', root, FOLDER)
    older = drive.add('spend-log-backup.json', oct_b, 'application/json', backup_body([entry('x3', 30)], '2026-10-21T12:00:00Z'))
    f1 = drive.add('Spend-Log-Backup.JSON', oct_a, 'application/json', backup_body([entry('x1', 10), entry('x2', 20)]))
    drive.add('SPEND-LOG-BACKUP.json', sep, 'application/json', backup_body([entry('s1', 5)]))
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()
    items = restore_list(d, 'Google Drive')
    assert [i.split('\n')[0] for i in items] == ['October 2026', 'September 2026'], items
    assert '3 records' in items[0] and '1 record' in items[1], items
    d.close()
    p = phone(b, srv, drive, entries=[entry('x9', 90)])   # backs up to Drive
    folders = [f['name'] for f in drive.live().values() if f['mime'] == FOLDER]
    assert sorted(folders) == ['OCTOBER_2026', 'September_2026', 'Spend-Expense-Backup', 'october_2026'], \
        'No new folders with the usual names: %s' % folders
    json_files = [(k, f) for k, f in drive.live().items() if f['name'].lower().endswith('.json') and f['parents'][0] in (oct_a, oct_b)]
    assert len(json_files) == 1 and json_files[0][0] == f1, 'Same-name copies combine into the newest file: %s' % \
        [(k, f['name']) for k, f in json_files]
    body = json.loads(drive.files[f1]['body'])
    got = {e['id']: e['amount'] for e in json.loads(body['data']['spendlog.v1'])['e']}
    assert got == {'x1': 10, 'x2': 20, 'x3': 30, 'x9': 90}, got
    assert drive.files[f1]['name'] == 'Spend-Log-Backup.JSON', 'The file keeps its own name'
    assert drive.files[older]['trashed'], 'The other copy goes to Drive trash'
    load_offer(p)
    assert amounts(p) == {'x1': 10, 'x2': 20, 'x3': 30, 'x9': 90}, amounts(p)
    p.close()


# ---- Safety -----------------------------------------------------------------
@test
def safety_compare_rules(b, srv):
    d = Device(b, srv).open()
    cases = [((0, 0), 'ok'), ((5, 0), 'ok'), ((0, 3), 'empty'), ((8, 10), 'ok'),
             ((4, 10), 'shrink'), ((2, 5), 'ok'), ((49, 100), 'shrink'), ((50, 100), 'ok'), ((120, 100), 'ok')]
    for (local, backup), want in cases:
        got = d.ev('SpendLogBackup.compareCounts(%d, %d)' % (local, backup))
        assert got == want, 'compare(%d, %d) = %s, expected %s' % (local, backup, got, want)
    d.close()


@test
def empty_device_never_overwrites_folder_backup(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    d.seed(10)
    d.backup_now()
    d.page.clock.set_system_time(datetime(2026, 10, 27, 19, 0))
    d.wipe_app_data()                                    # e.g. home-screen icon removed and added again
    assert 'restore' in d.banner().lower(), 'Should offer a restore, banner: ' + d.banner()
    d.run_auto()
    assert records_in(d.folder()[MONTH_FILE]) == 10, 'Automatic backup must not replace a good backup'
    d.backup_now()                                       # tap Back up now: merges, never empties the file
    assert records_in(d.folder()[MONTH_FILE]) == 10, 'Backup must still hold all 10 records'
    assert not d.dialogs, 'Merging needs no warning: %s' % d.dialogs
    assert '10 changes not on this device' in d.banner(), d.banner()
    d.page.click('#slb-banner')                          # load them back
    d.page.wait_for_timeout(1000)
    d.reload()
    assert d.records() == 10, 'Tapping the offer should bring the entries back'
    d.close()


@test
def empty_device_never_overwrites_drive_backup(b, srv):
    drive = FakeDrive()
    a = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()
    a.seed(8)
    a.page.click('#slb-fab')
    a.page.click('text=Back up to Google Drive')
    a.page.wait_for_timeout(600)
    a.close()
    fresh = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()
    fresh.page.click('#slb-fab')
    fresh.page.click('text=Back up to Google Drive')   # connecting Drive starts a backup (merge)
    fresh.page.wait_for_timeout(800)
    assert records_in(drive.tree()[MONTH_FILE]) == 8, 'Drive backup must be kept'
    assert not fresh.dialogs, fresh.dialogs
    assert '8 changes not on this device' in fresh.banner(), fresh.banner()
    fresh.page.click('#slb-banner')
    fresh.page.wait_for_timeout(1000)
    fresh.reload()
    assert fresh.records() == 8, 'New phone should get the entries from Drive'
    fresh.close()


@test
def much_smaller_data_is_held_until_confirmed(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    d.seed(20)
    d.backup_now()
    d.seed(6)                                            # lost most entries
    d.page.clock.set_system_time(datetime(2026, 10, 27, 19, 0))
    d.run_auto()
    assert records_in(d.folder()[MONTH_FILE]) == 20, 'Automatic run must hold the backup'
    assert 'paused' in d.banner().lower(), d.banner()
    d.answers = [True]                                   # person confirms on purpose
    d.backup_now()
    assert records_in(d.folder()[MONTH_FILE]) == 6
    d.close()


@test
def normal_deletes_still_back_up(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    d.seed(10)
    d.backup_now()
    d.seed(8)
    d.page.clock.set_system_time(datetime(2026, 10, 27, 19, 0))
    d.run_auto()
    assert records_in(d.folder()[MONTH_FILE]) == 8
    assert not d.dialogs, 'No warning expected for a few deletions'
    d.close()


@test
def empty_banner_goes_away_after_first_entry(b, srv):
    d = Device(b, srv, clock=OCT).open()
    d.page.clock.run_for(1000)
    assert 'restore' in d.banner().lower(), 'Fresh app should offer restore'
    d.add_expense(3, 'Snacks')
    d.page.clock.run_for(4000)
    d.page.wait_for_timeout(200)
    assert 'restore' not in d.banner().lower(), 'Restore offer should hide once entries exist'
    d.close()


# ---- Travel category ---------------------------------------------------------
CAPTURE_XLSX = """window.__xlsx = []; window.XLSX = { utils: {
  book_new: () => ({ SheetNames: [], Sheets: {} }),
  aoa_to_sheet: a => { __xlsx.push(a); return {}; }, json_to_sheet: a => { __xlsx.push(a); return {}; },
  book_append_sheet: () => {}, sheet_add_aoa: (ws, a) => { __xlsx.push(a); },
  encode_cell: () => 'A1', decode_range: () => ({ s: { r: 0, c: 0 }, e: { r: 0, c: 0 } }), encode_range: () => 'A1' },
  writeFile: () => {}, write: () => new Uint8Array() };"""


def set_entries(d, rows):
    """rows: (category id, amount, 'YYYY-MM-DD HH:MM', note). Saved exactly as the app saves them."""
    d.ev("""(rows) => {
      const e = rows.map((r, i) => { const [day, hm] = r[2].split(' '); const [y, mo, dd] = day.split('-').map(Number);
        const [h, mi] = hm.split(':').map(Number); const ts = new Date(y, mo - 1, dd, h, mi).getTime();
        return { id: 't' + i, amount: r[1], cat: r[0], sub: null, note: r[3] || '', ts: ts, source: 'manual',
                 createdAt: ts, split: null }; });
      localStorage.setItem('spendlog.v1', JSON.stringify({ e: e, m: [], p: [], s: [] }));
    }""", [list(r) for r in rows])
    d.reload()


@test
def travel_replaces_transport_in_the_app(b, srv):
    d = Device(b, srv)
    d.ctx.route('**/xlsx.full.min.js', lambda r: r.fulfill(status=200, content_type='text/javascript', body=CAPTURE_XLSX))
    d.open()
    d.page.click('#fab')
    d.page.wait_for_selector('#fAmt')
    labels = d.ev("[...document.querySelectorAll('#sheetRoot button')].map(b => b.textContent.trim())")
    assert 'Travel' in labels and 'Transport' not in labels, labels
    d.page.click('#closeSheet')
    d.add_expense(2.75, 'Travel')
    saved = d.ev("JSON.parse(localStorage.getItem('spendlog.v1')).e.map(x => x.cat)")
    assert saved == ['transport'], 'Travel must keep the saved id "transport" so old entries match: %s' % saved
    set_entries(d, [('transport', 18, '2026-10-20 17:45', 'Uber (saved before the rename)')])
    texts = []
    for view in ['recent', 'week', 'month']:
        d.page.click('button[data-view=%s]' % view)
        d.page.wait_for_timeout(200)
        texts.append(d.ev('document.body.innerText'))
    assert all('Travel' in t for t in texts[:1]), 'Older Transport entry should show as Travel'
    assert not any('Transport' in t for t in texts), 'The word Transport is still on screen'
    b64 = d.ev("""async () => { const a = new Uint8Array(await SpendLogBackup.reportBlob('all', 'All entries').arrayBuffer());
      let s = ''; for (let i = 0; i < a.length; i += 8192) s += String.fromCharCode.apply(null, a.subarray(i, i + 8192)); return btoa(s); }""")
    import io, base64, openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))
    sheets = json.dumps([str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value is not None])
    assert 'Travel' in sheets and 'Transport' not in sheets, 'Excel report should say Travel'
    d.close()


# ---- Backup arranged by category ------------------------------------------
MIXED = [
    ('lunch', 12.5, '2026-10-26 13:05', 'Chipotle'),
    ('transport', 2.75, '2026-10-26 08:40', 'bus'),
    ('lunch', 9, '2026-10-24 12:30', ''),
    ('grocery', 41.2, '2026-10-25 18:10', 'Harris Teeter'),
    ('breakfast', 4.5, '2026-10-26 08:05', 'coffee'),
    ('transport', 18, '2026-10-23 17:45', 'Uber'),
    ('grocery', 15, '2026-10-22 19:00', ''),
    ('lunch', 7.25, '2026-10-24 12:30', 'same time as another lunch'),
    ('gifts', 30, '2026-10-21 10:00', 'category the app may add later'),
]
APP_ORDER = ['Breakfast', 'Lunch', 'Dinner', 'Snacks', 'Grocery', 'Shopping', 'Household', 'Utilities', 'Travel', 'Other']


@test
def backup_is_grouped_by_category_and_sorted_by_time(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    set_entries(d, MIXED)
    d.backup_now()
    text = d.folder()[MONTH_FILE]
    backup = json.loads(text)
    groups = backup['byCategory']
    names = [g['category'] for g in groups]
    known = [n for n in names if n in APP_ORDER]
    assert known == sorted(known, key=APP_ORDER.index), 'Categories should follow the app order: %s' % names
    assert names == ['Breakfast', 'Lunch', 'Grocery', 'Travel', 'Gifts'], names
    assert 'Transport' not in names
    for g in groups:
        stamps = [(e['date'], e['time']) for e in g['entries']]
        assert stamps == sorted(stamps), '%s is not sorted by date and time: %s' % (g['category'], stamps)
        assert g['count'] == len(g['entries'])
        assert abs(g['total'] - sum(e['amount'] for e in g['entries'])) < 0.001
    lunch = next(g for g in groups if g['category'] == 'Lunch')
    assert [e['amount'] for e in lunch['entries']] == [9, 7.25, 12.5], 'Same-time entries keep the order they were added'
    assert lunch['group'] == 'Food' and lunch['total'] == 28.75
    travel = next(g for g in groups if g['category'] == 'Travel')
    assert [e.get('note') for e in travel['entries']] == ['Uber', 'bus']
    assert sum(g['count'] for g in groups) == len(MIXED), 'Every expense must appear exactly once'
    assert backup['records'] == len(MIXED)
    lines = [l.strip() for l in text.splitlines() if l.strip().startswith('{"date"')]
    assert len(lines) == len(MIXED), 'Each expense should be on its own line'
    assert text.index('"byCategory"') < text.index('"data"'), 'Readable part should come first'
    d.close()


@test
def backup_by_category_matches_on_drive_and_iphone(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    set_entries(d, MIXED[:4])
    d.page.click('#slb-fab')
    d.page.click('text=Back up to Google Drive')
    d.page.wait_for_timeout(800)
    on_drive = json.loads(d.drive.tree()[MONTH_FILE])
    assert [g['category'] for g in on_drive['byCategory']] == ['Lunch', 'Grocery', 'Travel']
    d.close()


@test
def empty_and_single_category_backups(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    assert json.loads(d.ev('SpendLogBackup.backupText()'))['byCategory'] == []
    set_entries(d, [('snacks', 3, '2026-10-26 15:00', '')])
    groups = json.loads(d.ev('SpendLogBackup.backupText()'))['byCategory']
    assert groups == [{'category': 'Snacks', 'group': 'Food', 'count': 1, 'total': 3,
                       'entries': [{'date': '2026-10-26', 'time': '15:00', 'amount': 3}]}], groups
    d.close()


@test
def restore_works_with_category_backup_and_older_backups(b, srv):
    d = Device(b, srv, kind='firefox', clock=OCT_EVENING).open()
    set_entries(d, MIXED)
    before = d.ev("localStorage.getItem('spendlog.v1')")
    new_style = d.ev('SpendLogBackup.backupText()')
    old_style = json.loads(new_style)
    del old_style['byCategory']                          # a backup made before this release
    for label, text in [('new', new_style), ('old', json.dumps(old_style))]:
        path = os.path.join(tempfile.mkdtemp(), 'spend-log-backup.json')
        open(path, 'w', encoding='utf-8').write(text)
        set_entries(d, [('other', 1, '2026-10-26 09:00', '')])
        d.ev('SpendLogBackup.restore()')
        with d.page.expect_file_chooser() as fc:
            d.page.click('#slb-modal button:has-text("Choose a file")')
        fc.value.set_files(path)
        d.page.wait_for_timeout(800)
        d.reload()
        assert d.ev("localStorage.getItem('spendlog.v1')") == before, label + ' backup did not restore exactly'
    d.close()


# ---- Excel copy of the backup, coloured like the app ----------------------
LABEL_TO_ID = {'Breakfast': 'breakfast', 'Lunch': 'lunch', 'Dinner': 'dinner', 'Snacks': 'snacks', 'Grocery': 'grocery',
               'Shopping': 'shopping', 'Household': 'household', 'Utilities': 'utilities', 'Travel': 'transport', 'Other': 'other'}


def app_colors():
    """Category colours from the app's own stylesheet (light theme = first definition)."""
    html = read('index.html')
    return {cid: re.search(r'--c-%s:#([0-9A-Fa-f]{6})' % cid, html).group(1).upper() for cid in LABEL_TO_ID.values()}


def open_xlsx(raw):
    data = base64.b64decode(raw[7:]) if isinstance(raw, str) and raw.startswith('base64:') else raw
    assert zipfile.ZipFile(io.BytesIO(data)).testzip() is None, 'Excel file is damaged'
    return openpyxl.load_workbook(io.BytesIO(data))


def fill_of(cell):
    f = cell.fill
    return (f.fgColor.rgb or '')[-6:].upper() if f is not None and f.fill_type == 'solid' else None


def read_category_sheet(ws):
    """[(label, band colour, band total, [(date, time, category, amount, note, row colour)])]"""
    out, cur = [], None
    for r in range(5, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if isinstance(b, str) and '·' in b:
            cur = (b.split('·')[0].strip(), fill_of(ws.cell(r, 2)), ws.cell(r, 5).value, [], fill_of(ws.cell(r, 1)))
            out.append(cur)
        elif cur and isinstance(b, datetime):
            cur[3].append((b, ws.cell(r, 3).value, ws.cell(r, 4).value, ws.cell(r, 5).value,
                           ws.cell(r, 6).value, fill_of(ws.cell(r, 4)), fill_of(ws.cell(r, 1))))
    return out


def lighter(light, strong):
    return all(int(light[i:i + 2], 16) > int(strong[i:i + 2], 16) for i in (0, 2, 4))


@test
def excel_backup_is_saved_and_coloured_like_the_app(b, srv):
    colors = app_colors()
    d = Device(b, srv, clock=OCT_EVENING).open()
    set_entries(d, [r for r in MIXED if r[0] != 'gifts'] + [('dinner', 23.4, '2026-10-25 20:15', ''),
                                                           ('utilities', 85, '2026-10-15 09:00', 'electric')])
    d.backup_now()
    files = d.folder()
    wb = open_xlsx(files[MONTH_XLSX])
    assert wb.sheetnames == ['By category', 'Summary', 'Money in'], wb.sheetnames
    groups_json = json.loads(files[MONTH_FILE])['byCategory']
    sheet = read_category_sheet(wb['By category'])
    assert [g[0] for g in sheet] == [g['category'] for g in groups_json], 'Excel and backup should list categories in the same order'
    grand = 0
    for label, band, total, rows, stripe in sheet:
        want = colors[LABEL_TO_ID[label]]
        assert band == want and stripe == want, '%s band is %s, app colour is %s' % (label, band, want)
        stamps = [(r[0], r[1]) for r in rows]
        assert stamps == sorted(stamps), label + ' rows are not sorted by date and time'
        assert all(r[2] == label for r in rows)
        assert all(isinstance(r[3], (int, float)) for r in rows), 'Amounts must be numbers so Excel can add them'
        assert abs(total - sum(r[3] for r in rows)) < 0.001, label + ' total is wrong'
        assert all(r[6] == want for r in rows), label + ' rows should have the colour stripe'
        assert all(r[5] and r[5] != want and lighter(r[5], want) for r in rows), label + ' rows should be a light shade'
        grand += total
    lunch = next(g for g in sheet if g[0] == 'Lunch')
    assert [(r[0].strftime('%Y-%m-%d'), r[1].strftime('%H:%M'), r[3]) for r in lunch[3]] == \
        [('2026-10-24', '12:30', 9), ('2026-10-24', '12:30', 7.25), ('2026-10-26', '13:05', 12.5)]
    summary = wb['Summary']
    rows = [(summary.cell(r, 2).value, summary.cell(r, 3).value, summary.cell(r, 4).value, fill_of(summary.cell(r, 1)))
            for r in range(4, summary.max_row + 1)]
    rows = rows[:[r[0] for r in rows].index('Total money out') + 1]          # the category part of the summary
    for (label, count, total, swatch), g in zip(rows, groups_json):
        assert (label, count, total) == (g['category'], g['count'], g['total']), (label, count, total, g)
        assert swatch == colors[LABEL_TO_ID[label]], label + ' summary colour does not match the app'
    assert rows[-1][0] == 'Total money out' and abs(rows[-1][2] - grand) < 0.001
    all_text = ' '.join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None)
    assert 'Transport' not in all_text and 'Travel' in all_text
    d.close()


@test
def excel_copy_goes_wherever_the_backup_goes(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()          # Google Drive
    set_entries(d, MIXED[:4])
    d.page.click('#slb-fab')
    d.page.click('text=Back up to Google Drive')
    d.page.wait_for_timeout(800)
    first = d.drive.tree()[MONTH_XLSX]
    assert [g[0] for g in read_category_sheet(open_xlsx(first)['By category'])] == ['Lunch', 'Grocery', 'Travel']
    d.add_expense(3, 'Snacks')
    d.backup_now()
    assert sorted(d.drive.tree()) == [MONTH_FILE, MONTH_XLSX], 'Excel copy on Drive should be replaced, not duplicated'
    assert 'Snacks' in [g[0] for g in read_category_sheet(open_xlsx(d.drive.tree()[MONTH_XLSX])['By category'])]
    d.close()
    f = Device(b, srv, kind='firefox', clock=OCT_EVENING).open()        # Downloads
    set_entries(f, MIXED[:3])
    got = []
    f.page.on('download', lambda dl: got.append(dl))
    f.ev('SpendLogBackup.backupNow()')
    f.page.wait_for_timeout(1500)
    names = sorted(dl.suggested_filename for dl in got)
    assert names == ['spend-expense-backup_October_2026.json', 'spend-expense-backup_October_2026.xlsx'], names
    xl = next(dl for dl in got if dl.suggested_filename.endswith('.xlsx'))
    open_xlsx(open(xl.path(), 'rb').read())
    f.close()


@test
def excel_handles_special_text_and_empty_backups(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    d.backup_now()
    files = d.folder()
    ws = open_xlsx(files[MONTH_XLSX])['By category']
    assert any(c.value == 'No expenses yet' for row in ws.iter_rows() for c in row), 'Empty backup should say so'
    note = 'Tom & Jerry\'s <deli> "special" — café ☕'
    set_entries(d, [('lunch', 1234.5, '2026-10-26 12:00', note)])
    d.backup_now()
    ws = open_xlsx(d.folder()[MONTH_XLSX])['By category']
    rows = read_category_sheet(ws)[0][3]
    assert rows[0][4] == note, 'Note text changed: %r' % rows[0][4]
    assert rows[0][3] == 1234.5 and '"$"#,##0.00' in ws.cell(6, 5).number_format.replace('\\', '')
    d.close()


@test
def excel_problem_never_blocks_the_main_backup(b, srv):
    d = Device(b, srv, clock=OCT_EVENING)
    d.ctx.add_init_script("""(() => { const B = window.Blob; window.Blob = function (parts, opts) {
        if (opts && /spreadsheetml/.test(opts.type || '')) throw new Error('simulated Excel failure');
        return new B(parts, opts); }; window.Blob.prototype = B.prototype; })();""")
    d.open()
    set_entries(d, MIXED[:3])
    d.backup_now()
    files = d.folder()
    assert MONTH_FILE in files and records_in(files[MONTH_FILE]) == 3, 'Main backup must still be saved'
    assert MONTH_XLSX not in files
    assert d.meta()['last'] == '2026-10-26'
    d.close()


# ---- Font: Plus Jakarta Sans, stored with the app --------------------------
FONT = 'Plus Jakarta Sans'


def font_files_in_page():
    return re.findall(r'url\((fonts/[^)]+\.woff2)\)', read('index.html'))


@test
def font_is_stored_with_the_app_and_used_everywhere(b, srv):
    html, sw = read('index.html'), read('sw.js')
    assert 'fonts.googleapis.com' not in html and 'fonts.gstatic.com' not in html, 'Google Fonts should be gone'
    refs = font_files_in_page()
    assert len(refs) == 10, 'Expected 5 weights x 2 character sets, got %d' % len(refs)
    for f in refs:
        assert os.path.getsize(os.path.join(APP_DIR, f)) > 5000, f + ' is missing'
        assert "'/" + f + "'" in sw, f + ' must be kept for offline use (sw.js)'
    assert os.path.exists(os.path.join(APP_DIR, 'fonts', 'OFL-PlusJakartaSans.txt')), 'Font licence should ship with the fonts'
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING)
    outside = []
    d.page.on('request', lambda r: outside.append(r.url) if 'fonts.g' in r.url else None)
    d.open()
    set_split_state(d, FRIENDS, SPLITS, PAYMENTS)
    d.ev('document.fonts.ready.then(() => 1)')
    assert d.ev("document.fonts.check('800 22px \"%s\"') && document.fonts.check('400 16px \"%s\"')" % (FONT, FONT))
    loaded = d.ev("[...document.fonts].filter(f => f.family.includes('Jakarta') && f.status === 'loaded').map(f => f.weight)")
    assert {'400', '800'} <= set(loaded), 'Font files did not load: %s' % loaded
    fam = lambda sel: d.ev("(s) => getComputedStyle(document.querySelector(s)).fontFamily", sel)
    h1 = d.ev("(() => { const s = getComputedStyle(document.querySelector('h1')); return [s.fontFamily, s.fontWeight]; })()")
    assert h1[0].startswith('"%s"' % FONT) and h1[1] == '800', h1
    for sel in ['body', '.big', '#slb-fab', '#exportBtn', 'nav.tabs button']:
        assert fam(sel).startswith('"%s"' % FONT), '%s uses %s' % (sel, fam(sel))
    d.page.click('#slb-fab')
    for sel in ['#slb-panel', '#slb-panel button', '#slb-status']:
        assert '"%s"' % FONT in fam(sel), '%s uses %s' % (sel, fam(sel))
    d.page.click('body', position={'x': 5, 'y': 600})
    open_summary(d, 'Ravi')
    for sel in ['#waSend', '#sumText', '#waPhone']:
        assert fam(sel).startswith('"%s"' % FONT), '%s uses %s' % (sel, fam(sel))
    assert d.ev("getComputedStyle(document.querySelector('.big')).fontVariantNumeric") == 'tabular-nums', 'Totals should use aligned digits'
    assert d.ev("getComputedStyle(document.querySelector('.num')).fontVariantNumeric") == 'tabular-nums'
    assert not outside, 'No requests to Google Fonts: %s' % outside
    d.page.goto(srv.url + '/privacy.html')
    d.ev('document.fonts.ready.then(() => 1)')
    assert fam('body').startswith('"%s"' % FONT)
    d.ctx.close()


@test
def font_works_offline(b, srv):
    d = Device(b, srv, kind='iphone', sw=True)
    d.page.goto(srv.url + '/')
    d.page.wait_for_timeout(2500)                      # offline copy (sw.js) saves the app and fonts
    d.reload()
    d.page.wait_for_timeout(500)
    cached = d.ev("caches.keys().then(k => caches.open(k[0])).then(c => c.keys()).then(r => r.map(x => new URL(x.url).pathname))")
    for f in font_files_in_page():
        assert '/' + f in cached, f + ' not saved for offline use'
    d.ctx.set_offline(True)
    d.reload()
    ok = d.ev("""document.fonts.load('800 22px "%s"').then(f => f.length > 0 && f.every(x => x.status === 'loaded'))""" % FONT)
    assert ok, 'Font should load offline'
    assert d.ev("getComputedStyle(document.querySelector('h1')).fontFamily").startswith('"%s"' % FONT)
    d.ctx.set_offline(False)
    d.ctx.close()


# ---- Merge instead of replace (several devices, one Google Drive) ---------
def entry(i, amount, day='2026-10-26 12:00', cat='lunch', note=''):
    ts = int(datetime.strptime(day, '%Y-%m-%d %H:%M').timestamp() * 1000)
    return {'id': i, 'amount': amount, 'cat': cat, 'sub': None, 'note': note, 'ts': ts,
            'source': 'manual', 'createdAt': ts, 'split': None}


def change(d, *ops):
    """Change this device's saved entries: ('add', entry) | ('edit', id, amount) | ('del', id). Then reopen the app."""
    d.ev("""(ops) => {
      const o = JSON.parse(localStorage.getItem('spendlog.v1') || '{"e":[],"m":[],"p":[],"s":[]}');
      for (const op of ops) {
        if (op[0] === 'add') o.e.push(op[1]);
        if (op[0] === 'edit') o.e.find(x => x.id === op[1]).amount = op[2];
        if (op[0] === 'del') o.e = o.e.filter(x => x.id !== op[1]);
      }
      localStorage.setItem('spendlog.v1', JSON.stringify(o));
    }""", [list(o) for o in ops])
    d.reload()


def amounts(d):
    """{id: amount} on this device."""
    o = json.loads(d.ev("localStorage.getItem('spendlog.v1') || '{}'"))
    return {e['id']: e['amount'] for e in o.get('e', [])}


def drive_amounts(drive):
    b = json.loads(drive.tree()[MONTH_FILE])
    return {e['id']: e['amount'] for e in json.loads(b['data']['spendlog.v1'])['e']}


def backup_body(entries, created='2026-10-20T12:00:00Z', deleted=None, fmt=2):
    b = {'app': 'Spenderrr', 'format': fmt, 'createdAt': created,
         'data': {'spendlog.v1': json.dumps({'e': entries, 'm': [], 'p': [], 's': []})}}
    if deleted is not None:
        b['deleted'] = deleted
    return json.dumps(b)


def load_offer(d):
    """Tap 'changes not on this device' and wait for the app to reload."""
    assert 'not on this device' in d.banner(), 'Expected an offer to load changes, banner: %r' % d.banner()
    d.page.click('#slb-banner')
    d.page.wait_for_timeout(1000)
    d.reload()


def phone(b, srv, drive, entries=(), paused=True):
    """A phone using the shared Google Drive, with Drive connected and its clock paused (no surprise timers)."""
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()
    if entries:
        change(d, *[('add', e) for e in entries])
    d.page.click('#slb-fab')
    d.page.click('text=Back up to Google Drive')
    d.page.wait_for_timeout(800)
    if d.page.is_visible('#slb-panel'):
        d.page.click('h1')                                # tap outside to close the Backup panel
    if paused:
        d.page.clock.pause_at(datetime(2026, 10, 26, 19, 30))
    return d


@test
def merge_rules(b, srv):
    d = Device(b, srv).open()
    run = lambda local, remote, base, deleted=None, now=10 ** 13: d.ev(
        """([l, r, b, del, now]) => SpendLogBackup.mergeStores({ e: l }, { e: r }, b ? SpendLogBackup.fingerprints({ e: b }) : null, del, now)""",
        [local, remote, base, deleted, now])
    ids = lambda res: sorted(x['id'] for x in res['store']['e'])
    a, a2, bb, c = entry('a', 1), entry('a', 2), entry('b', 5), entry('c', 9)
    assert ids(run([a], [bb], None)) == ['a', 'b'], 'first save: everything from both sides'
    assert run([a2], [a], [a])['store']['e'] == [a2], 'edited here -> this device wins'
    assert run([a], [a2], [a])['store']['e'] == [a2], 'edited elsewhere -> take the other device'
    assert run([a2], [entry('a', 3)], [a])['store']['e'] == [a2], 'edited on both -> the device saving now wins'
    res = run([], [a, bb], [a, bb])
    assert res['store']['e'] == [] and set(res['deleted']['e']) == {'a', 'b'}, 'deleted here -> recorded'
    assert run([a], [], [a], {'e': {'a': 1}})['store']['e'] == [], 'deleted elsewhere, unchanged here -> gone'
    assert run([a2], [], [a], {'e': {'a': 1}})['store']['e'] == [a2], 'deleted elsewhere but edited here -> kept'
    assert run([a, c], [a], [a, c])['store']['e'] == [a, c], 'missing without a recorded deletion -> never lost'
    assert run([], [bb], [], {'e': {'b': 1}})['store']['e'] == [], 'stale copy of a deleted entry is dropped'
    old = run([], [a], [a], None, 10 ** 13)['deleted']
    assert 'a' in old['e'] and run([], [], None, old, 10 ** 13 + 181 * 86400000)['deleted'] == {}, 'old deletions are forgotten'
    people = d.ev("""SpendLogBackup.mergeStores({ p: [{ name: 'Ravi', addedAt: 1 }] }, { p: [{ name: 'Ravi', addedAt: 2 }, { name: 'Priya', addedAt: 3 }] }, null, null)""")
    assert [p['name'] for p in people['store']['p']] == ['Ravi', 'Priya'], 'people are matched by name'
    d.close()


@test
def two_phones_share_one_drive_without_losing_entries(b, srv):
    drive = FakeDrive()
    a = phone(b, srv, drive, [entry('a1', 10), entry('a2', 20), entry('a3', 30)])
    assert drive_amounts(drive) == {'a1': 10, 'a2': 20, 'a3': 30}
    bp = phone(b, srv, drive, [entry('b1', 5), entry('b2', 6)])
    assert drive_amounts(drive) == {'a1': 10, 'a2': 20, 'a3': 30, 'b1': 5, 'b2': 6}, "Phone B must not wipe phone A's entries"
    assert '3 changes not on this device' in bp.banner(), bp.banner()
    load_offer(bp)
    assert amounts(bp) == drive_amounts(drive)
    a.backup_now()
    assert '2 changes not on this device' in a.banner(), a.banner()
    load_offer(a)
    assert amounts(a) == amounts(bp) == drive_amounts(drive), 'Both phones should end up the same'
    assert len(drive.copies('spend-log-backup.json')) == 1, 'Still one backup file'
    a.close(); bp.close()


@test
def edits_and_deletions_reach_the_other_phone(b, srv):
    drive = FakeDrive()
    a = phone(b, srv, drive, [entry('a1', 10), entry('a2', 20), entry('a3', 30)])
    bp = phone(b, srv, drive, [entry('b1', 5)])
    load_offer(bp)
    a.backup_now(); load_offer(a)
    change(a, ('del', 'a1'), ('edit', 'b1', 7))           # phone A deletes one, edits one
    a.backup_now()
    assert drive_amounts(drive) == {'a2': 20, 'a3': 30, 'b1': 7}
    bp.backup_now()                                        # phone B hasn't changed anything
    assert '2 changes not on this device' in bp.banner(), bp.banner()
    load_offer(bp)
    assert amounts(bp) == {'a2': 20, 'a3': 30, 'b1': 7}, 'Deletion and edit should reach phone B'
    change(a, ('edit', 'a2', 50))                          # both edit the same entry: last save wins
    change(bp, ('edit', 'a2', 70))
    a.backup_now(); bp.backup_now()
    assert drive_amounts(drive)['a2'] == 70
    change(a, ('del', 'a3'))                               # A deletes while B edits it: the edit is kept
    change(bp, ('edit', 'a3', 33))
    a.backup_now(); bp.backup_now()
    assert drive_amounts(drive)['a3'] == 33, 'An edit made after a deletion elsewhere should be kept'
    a.backup_now(); load_offer(a)
    assert amounts(a) == amounts(bp) == drive_amounts(drive) == {'a2': 70, 'a3': 33, 'b1': 7}
    a.close(); bp.close()


@test
def older_app_version_overwriting_the_file_loses_nothing(b, srv):
    drive = FakeDrive()
    a = phone(b, srv, drive, [entry('e%d' % i, i + 1) for i in range(5)])
    fid = drive.copies('spend-log-backup.json')[0][0]
    drive.files[fid]['body'] = backup_body([entry('e0', 1), entry('e1', 2)], fmt=1)   # old app replaced the file
    a.backup_now()
    assert drive_amounts(drive) == {'e%d' % i: i + 1 for i in range(5)}, 'Entries missing from the file must come back'
    assert amounts(a) == drive_amounts(drive) and 'not on this device' not in a.banner()
    a.close()


@test
def backs_up_after_each_change_and_waits_when_signed_out(b, srv):
    drive = FakeDrive()
    d = Device(b, srv, kind='iphone', clock=OCT, drive=drive).open()   # 5 PM
    d.seed(2)
    connect_drive(d)
    d.add_expense(7, 'Snacks')
    d.page.clock.run_for(5000)                            # 3 seconds after the change
    d.page.wait_for_timeout(800)
    assert len(drive_amounts(drive)) == 3, 'New entry should reach Drive without a tap'
    assert not d.meta().get('pending')
    d.page.clock.run_for(2 * 3600 * 1000)                 # sign-in expires (about an hour)
    d.add_expense(8, 'Snacks')
    d.page.clock.run_for(5000)
    d.page.wait_for_timeout(500)
    assert len(drive_amounts(drive)) == 3, 'Without sign-in nothing is sent'
    assert d.meta().get('pending') is True
    d.page.click('#slb-fab')
    assert 'Changes waiting' in d.page.inner_text('#slb-status')
    d.page.clock.run_for(61000)                           # it's after 6 PM: daily check asks for a tap
    d.page.wait_for_timeout(300)
    assert 'Google Drive' in d.banner(), d.banner()
    d.page.click('#slb-banner')
    d.page.wait_for_timeout(800)
    assert len(drive_amounts(drive)) == 4 and not d.meta().get('pending')
    d.close()
    f = Device(b, srv, kind='android', clock=OCT).open()  # Android folder: also after each change
    f.seed(2)
    f.backup_now()
    f.add_expense(9, 'Lunch')
    f.page.clock.run_for(5000)
    f.page.wait_for_timeout(800)
    assert records_in(f.folder()[MONTH_FILE]) == 3, 'Folder backup should follow each change'
    f.close()


@test
def drive_duplicates_are_merged_into_the_newest_file(b, srv):
    drive = FakeDrive()
    r1 = drive.add('spend-expense-backup', 'root', FOLDER)
    r2 = drive.add('spend-expense-backup', 'root', FOLDER)            # duplicate folders
    m1 = drive.add('October_2026', r1, FOLDER)
    m2 = drive.add('October_2026', r2, FOLDER)
    older = drive.add('spend-log-backup.json', m1, 'application/json', backup_body([entry('x1', 10), entry('x2', 20)]))
    newer = drive.add('spend-log-backup.json', m2, 'application/json', backup_body([entry('x2', 25), entry('x3', 30)]))
    d = phone(b, srv, drive)
    left = drive.copies('spend-log-backup.json')
    assert [i for i, _ in left] == [newer], 'Only the newest file should remain: %s' % [i for i, _ in left]
    assert drive.files[older]['trashed'], 'The duplicate should be in Drive trash (recoverable), not deleted'
    assert len(drive.copies('spend-log-backup.xlsx')) == 1
    body = json.loads(left[0][1])
    assert {e['id']: e['amount'] for e in json.loads(body['data']['spendlog.v1'])['e']} == {'x1': 10, 'x2': 25, 'x3': 30}, \
        'Duplicates combine, newest copy wins'
    assert 'moved to Drive trash' in d.page.inner_text('#slb-toast'), d.page.inner_text('#slb-toast')
    load_offer(d)
    assert amounts(d) == {'x1': 10, 'x2': 25, 'x3': 30}
    d.close()


@test
def restore_list_shows_every_backup(b, srv):
    drive = FakeDrive()
    root = drive.add('spend-expense-backup', 'root', FOLDER)
    sep = drive.add('September_2026', root, FOLDER)
    octo = drive.add('October_2026', root, FOLDER)
    drive.add('Notes', root, FOLDER)                                     # not a month: ignored
    drive.add('spend-log-backup.json', sep, 'application/json', backup_body([entry('s%d' % i, i) for i in range(3)], '2026-09-30T23:00:00Z'))
    drive.add('spend-log-backup.json', octo, 'application/json', backup_body([entry('o%d' % i, i) for i in range(5)], '2026-10-26T23:00:00Z'))
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()
    d.ev('SpendLogBackup.restore()')
    d.page.click('#slb-modal button:has-text("Google Drive")')
    d.page.wait_for_selector('#slb-modal .slb-item')
    items = [d.page.locator('#slb-modal .slb-item').nth(i).inner_text() for i in range(d.page.locator('#slb-modal .slb-item').count())]
    assert len(items) == 2, items
    assert items[0].startswith('October 2026') and '5 records' in items[0], items
    assert items[1].startswith('September 2026') and '3 records' in items[1], items
    d.page.locator('#slb-modal .slb-item').nth(1).click()
    d.page.wait_for_timeout(800)
    d.reload()
    assert sorted(amounts(d)) == ['s0', 's1', 's2'], 'Should restore the September backup'
    d.close()
    e = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=FakeDrive()).open()   # nothing in Drive
    e.ev('SpendLogBackup.restore()')
    e.page.click('#slb-modal button:has-text("Google Drive")')
    for _ in range(50):
        if 'No backups' in e.page.inner_text('#slb-list'):
            break
        e.page.wait_for_timeout(100)
    assert 'No backups found in Google Drive (saberi.test@gmail.com).' in e.page.inner_text('#slb-list')
    e.close()
    f = Device(b, srv, kind='android', clock=datetime(2026, 9, 30, 19, 0)).open()     # folder: every month too
    f.seed(2)
    f.backup_now()
    f.page.clock.set_system_time(datetime(2026, 10, 5, 19, 0))
    f.add_expense(4, 'Lunch')
    f.backup_now()
    f.ev('SpendLogBackup.restore()')
    f.page.click('#slb-modal button:has-text("Backup folder")')
    f.page.wait_for_selector('#slb-modal .slb-item')
    got = [f.page.locator('#slb-modal .slb-item').nth(i).inner_text() for i in range(f.page.locator('#slb-modal .slb-item').count())]
    assert [g.split('\n')[0] for g in got] == ['October 2026', 'September 2026'], got
    assert '3 records' in got[0] and '2 records' in got[1], got
    f.close()


# ---- Which Google account Drive uses -------------------------------------
def connect_drive(d):
    d.page.click('#slb-fab')
    d.page.click('text=Back up to Google Drive')
    d.page.wait_for_timeout(800)


def account_card(d):
    el = d.page.query_selector('#slb-account')
    return el.inner_text() if el else ''


@test
def drive_shows_the_connected_google_account(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    d.seed(2)
    connect_drive(d)
    assert 'saberi.test@gmail.com' in d.page.inner_text('#slb-toast'), 'Connect message should name the account'
    card = account_card(d)
    assert 'Saberi R' in card and 'saberi.test@gmail.com' in card, card
    assert 'Google Drive (saberi.test@gmail.com)' in d.page.inner_text('#slb-status')
    assert any('/drive/v3/about' in c for c in d.drive.calls)
    backup = d.drive.tree()[MONTH_FILE]
    assert 'saberi.test@gmail.com' not in backup and 'Saberi R' not in backup, 'Account details must not go into backups'
    d.ev('sessionStorage.clear()')                     # app reopened later: no sign-in yet
    d.reload()
    d.page.click('#slb-fab')
    assert 'saberi.test@gmail.com' in account_card(d), 'Account should still show after reopening'
    d.page.click('#slb-panel button:has-text("Disconnect Google Drive")')
    d.page.wait_for_timeout(300)
    assert 'saberi.test@gmail.com' in d.page.inner_text('#slb-toast')
    assert not d.page.query_selector('#slb-account'), 'Account card should go away after disconnecting'
    assert not d.meta().get('driveAccount'), 'Account details should be removed on disconnect'
    d.close()


@test
def drive_account_switch_and_lookup_problems(b, srv):
    drive = FakeDrive()
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=drive).open()
    d.seed(3)
    connect_drive(d)
    drive.account = {'displayName': 'Saberi Work', 'emailAddress': 'saberi.work@example.com'}
    d.ev('sessionStorage.clear()')
    d.reload()
    d.backup_now()                                      # signs in again, with the other account
    assert 'now Saberi Work (saberi.work@example.com)' in d.page.inner_text('#slb-toast'), d.page.inner_text('#slb-toast')
    assert 'was saberi.test@gmail.com' in d.page.inner_text('#slb-toast')
    d.page.click('#slb-fab')
    assert 'saberi.work@example.com' in account_card(d)
    d.close()
    other = FakeDrive(); other.account = None              # Google refuses the account lookup
    e = Device(b, srv, kind='iphone', clock=OCT_EVENING, drive=other).open()
    e.seed(2)
    connect_drive(e)
    assert records_in(other.tree()[MONTH_FILE]) == 2, 'Backup must still work without the account name'
    assert 'account name shows after the next Drive backup' in account_card(e), account_card(e)
    e.close()
    f = Device(b, srv, kind='iphone', clock=OCT_EVENING,
               drive=FakeDrive({'displayName': 'Only Name'})).open()   # Google hides the email
    f.seed(1)
    connect_drive(f)
    assert 'Only Name' in account_card(f) and 'Google Drive (Only Name)' in f.page.inner_text('#slb-status')
    f.close()


# ---- Send on WhatsApp ------------------------------------------------------
def set_split_state(d, people, entries, payments=()):
    """entries: (cat, bill, 'YYYY-MM-DD HH:MM', note, paid_by, {name: share}); payments: (name, amt, 'in'|'out', when)."""
    d.ev("""(a) => {
      const t = s => { const [day, hm] = s.split(' '); const [y, m, dd] = day.split('-').map(Number);
        const [h, mi] = hm.split(':').map(Number); return new Date(y, m - 1, dd, h, mi).getTime(); };
      const e = a.entries.map((r, i) => ({ id: 'w' + i, amount: r[1], cat: r[0], sub: null, note: r[3], ts: t(r[2]),
        source: 'manual', createdAt: t(r[2]), split: { paidBy: r[4], mode: 'equal',
        parts: Object.entries(r[5]).map(([name, amt]) => ({ name, amt })), my: r[5].Me || 0 } }));
      const s = a.payments.map((p, i) => ({ id: 'p' + i, name: p[0], amt: p[1], dir: p[2], ts: t(p[3]) }));
      const p = a.people.map((name, i) => ({ name, addedAt: 1000 + i }));
      localStorage.setItem('spendlog.v1', JSON.stringify({ e, m: [], p, s }));
    }""", {'people': list(people), 'entries': [list(x) for x in entries], 'payments': [list(x) for x in payments]})
    d.reload()


FRIENDS = ['Ravi', 'Priya', 'Asha', 'Ken']
SPLITS = [
    ('dinner', 60, '2026-10-25 20:15', 'Thai Palace', 'Me', {'Me': 20, 'Ravi': 20, 'Priya': 20}),
    ('transport', 30, '2026-10-26 08:40', 'Uber to airport', 'Me', {'Me': 15, 'Ravi': 15}),
    ('lunch', 24, '2026-10-24 12:30', 'Pizza', 'Ravi', {'Me': 12, 'Ravi': 12}),
    ('snacks', 9, '2026-10-23 16:00', 'Chips', 'Priya', {'Priya': 3, 'Ravi': 3, 'Me': 3}),
    ('grocery', 50, '2026-10-22 18:00', 'Costco run', 'Asha', {'Me': 25, 'Asha': 25}),
    ('utilities', 20, '2026-10-20 09:00', 'Wi-Fi', 'Me', {'Me': 10, 'Ken': 10}),
]
PAYMENTS = [('Ravi', 10, 'in', '2026-10-26 09:00'), ('Ken', 10, 'in', '2026-10-21 10:00')]

RAVI_MESSAGE = """🧾 *Split summary · Ravi*
_As of Mon 26 Oct 2026_

*Ravi owes you*
🍽️ Dinner · Thai Palace — *$20.00*
      Sun 25 Oct · bill $60.00 · Ravi's share
🚕 Travel · Uber to airport — *$15.00*
      Mon 26 Oct · bill $30.00 · Ravi's share

*You owe Ravi*
🥪 Lunch · Pizza — *$12.00*
      Sat 24 Oct · bill $24.00 · your share

*Payments*
✅ Ravi paid you $10.00 · Mon 26 Oct

━━━━━━━━━━━━━━
💰 *Total: Ravi owes you $13.00*

_Sent from Spenderrr_"""


def open_summary(d, name):
    d.page.click('button[data-view=split]')
    d.page.click('main .t1:text-is("%s")' % name)
    d.page.wait_for_selector('#fSummary')
    d.page.click('#fSummary')
    d.page.wait_for_selector('#waSend')
    return d.ev("document.getElementById('sumText').value")


def close_sheet(d):
    d.page.click('#closeSheet')
    d.page.wait_for_timeout(200)


@test
def whatsapp_message_is_structured_and_correct(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    set_split_state(d, FRIENDS, SPLITS, PAYMENTS)
    text = open_summary(d, 'Ravi')
    assert text == RAVI_MESSAGE, 'Ravi message differs:\n' + text
    assert 'Spend Log' not in text and 'Transport' not in text
    close_sheet(d)
    asha = open_summary(d, 'Asha')
    assert '*You owe Asha*' in asha and '*Asha owes you*' not in asha and '💸 *Total: You owe Asha $25.00*' in asha, asha
    close_sheet(d)
    ken = open_summary(d, 'Ken')
    assert '✅ *All settled up*' in ken and '✅ Ken paid you $10.00 · Wed 21 Oct' in ken, ken
    close_sheet(d)
    priya = open_summary(d, 'Priya')
    assert 'Chips' in priya and 'Total: Priya owes you $17.00' in priya, priya
    assert 'Ravi' not in priya.replace('Split summary · Priya', ''), "Other friends' debts must not appear"
    d.close()


RAVI_MESSAGE_NAMED = RAVI_MESSAGE.replace('Split summary · Ravi*', 'Split summary · Ravi & Saberi*') \
    .replace('Ravi owes you', 'Ravi owes Saberi').replace('You owe Ravi', 'Saberi owes Ravi') \
    .replace('your share', "Saberi's share").replace('Ravi paid you', 'Ravi paid Saberi')


@test
def whatsapp_message_uses_names_not_i_and_you(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    set_split_state(d, FRIENDS + ['Saberi'], SPLITS + [('dinner', 40, '2026-10-26 18:00', 'Biryani', 'Me', {'Me': 20, 'Saberi': 20})], PAYMENTS)
    connect_drive(d)                                   # Google account name: "Saberi R"
    d.page.click('h1')
    text = open_summary(d, 'Ravi')
    assert text == RAVI_MESSAGE_NAMED, 'Named Ravi message differs:\n' + repr(text)
    for word in ('You ', 'you ', ' me ', 'I owe', ' my '):
        assert word not in text.replace('_Sent from Spenderrr_', ''), 'Message should not say "%s"' % word.strip()
    close_sheet(d)
    asha = open_summary(d, 'Asha')
    assert '*Saberi owes Asha*' in asha and '💸 *Total: Saberi owes Asha $25.00*' in asha, asha
    close_sheet(d)
    ken = open_summary(d, 'Ken')
    assert '✅ Ken paid Saberi $10.00 · Wed 21 Oct' in ken, ken
    close_sheet(d)
    same = open_summary(d, 'Saberi')                   # friend with the same first name: use the full name
    assert '*Saberi owes Saberi R*' in same and "Saberi's share" in same, same
    close_sheet(d)
    # Odd account names are cleaned; an email-only account falls back to "you"
    names = d.ev("""() => ['Sa*be_ri <b>X', 'only@email.com', '  '].map(n => {
        const m = JSON.parse(localStorage.getItem('__spendlog_backup_meta') || '{}'); m.driveAccount = { name: n, email: 'a@b.c' };
        localStorage.setItem('__spendlog_backup_meta', JSON.stringify(m)); return SpenderrrWhatsApp.senderName('Ravi'); })""")
    assert names == ['Saberi', '', ''], names
    d.ev("""() => { const m = JSON.parse(localStorage.getItem('__spendlog_backup_meta')); m.driveAccount = null;
        localStorage.setItem('__spendlog_backup_meta', JSON.stringify(m)); }""")
    assert open_summary(d, 'Ravi') == RAVI_MESSAGE, 'Without a Drive name the message says "you"'
    d.close()


@test
def whatsapp_totals_match_the_app_for_every_friend(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    set_split_state(d, FRIENDS, SPLITS, PAYMENTS)
    d.page.click('button[data-view=split]')
    text = d.ev("document.querySelector('main').innerText")
    for name in FRIENDS:
        m = re.search(r'(?m)^%s\n(.+)\n(\$[\d,]+\.\d\d|—)$' % name, text)
        assert m, 'Could not find %s on the Split screen' % name
        status, amount = m.group(1), m.group(2)
        bal = d.ev('(n) => SpenderrrWhatsApp.ledger(n).balance', name)
        if bal == 0:
            assert 'Settled' in status, '%s: message says settled, app says %r %s' % (name, status, amount)
        else:
            assert amount == '$%.2f' % abs(bal), '%s: message total %.2f, app shows %s' % (name, bal, amount)
            assert ('owes you' in status) == (bal > 0), '%s direction differs from app: %r' % (name, status)
    d.close()


@test
def send_on_whatsapp_opens_chat_with_the_message(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING)
    d.ctx.route('https://wa.me/**', lambda r: r.fulfill(status=200, body='WhatsApp'))
    d.ctx.add_init_script("""window.__sent = [];
      Object.defineProperty(navigator, 'clipboard', { configurable: true, value: { writeText: async t => __sent.push(['copy', t]) } });
      const realShare = navigator.share;
      Object.defineProperty(navigator, 'share', { configurable: true, value: async x => {
        if (x && x.text) __sent.push(['share', x.text]); else if (realShare) return realShare.call(navigator, x); } });""")
    d.open()
    set_split_state(d, FRIENDS, SPLITS, PAYMENTS)
    text = open_summary(d, 'Ravi')
    with d.ctx.expect_page() as pop:
        d.page.click('#waSend')
    url = pop.value.url
    assert url.startswith('https://wa.me/?text='), 'Without a number WhatsApp should let you pick the chat: ' + url
    assert unquote(url.split('text=', 1)[1]) == text, 'Link text must equal the summary'
    pop.value.close()
    d.page.fill('#waPhone', '919 555 1234')
    d.page.click('#waSave')
    assert 'country code' in d.page.inner_text('#waNote')
    assert d.ev("localStorage.getItem('spenderrr.contacts')") in (None, '{}'), 'Invalid number must not be saved'
    d.page.fill('#waPhone', '+1 (919) 555-1234')
    d.page.click('#waSave')
    assert '+19195551234' in d.page.inner_text('#waNote')
    with d.ctx.expect_page() as pop:
        d.page.click('#waSend')
    assert pop.value.url.startswith('https://wa.me/19195551234?text='), pop.value.url
    pop.value.close()
    d.page.click('#copySum')
    d.page.click('#shareSum')
    d.page.wait_for_timeout(300)
    assert d.ev('window.__sent') == [['copy', text], ['share', text]], 'Copy and Share should send the new format'
    close_sheet(d)
    d.reload()
    open_summary(d, 'Ravi')
    assert d.page.input_value('#waPhone') == '+19195551234', 'Saved number should be remembered'
    assert json.loads(json.loads(d.ev('SpendLogBackup.backupText()'))['data']['spenderrr.contacts']) == {'Ravi': '+19195551234'}, \
        'Saved numbers should be in the backup'
    for raw, want in [('0091 98765 43210', '+919876543210'), ('+44 7700 900123', '+447700900123'), ('+1 555', None),
                      ('9876543210', None), ('', '')]:
        assert d.ev('(r) => SpenderrrWhatsApp.normalizePhone(r)', raw) == want, raw
    d.close()


@test
def whatsapp_message_handles_special_text_and_long_history(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    many = [('lunch', 10, '2026-10-%02d 12:%02d' % (1 + i // 2, (i % 2) * 30), 'Meal %d' % (i + 1), 'Me', {'Me': 5, 'Ravi': 5})
            for i in range(30)]
    tricky = ('dinner', 40, '2026-10-26 18:00', 'Tom & Jerry\'s *best* _pasta_ ~50% off~ #1 ?=&', 'Me', {'Me': 20, 'Ravi': 20})
    set_split_state(d, ['Ravi'], many + [tricky])
    text = open_summary(d, 'Ravi')
    assert "Tom & Jerry's best pasta 50% off #1 ?=&" in text, 'Formatting characters should be removed from notes'
    assert text.count('— *') == 25, 'Long lists should show the latest 25 items'
    assert '_…6 earlier items included in the total_' in text, text[:300]
    assert 'Total: Ravi owes you $170.00' in text, 'Total must include every item'
    assert 'Meal 30' in text and 'Meal 6\n' not in text and 'Meal 7 ' in text
    url = d.ev("(t) => SpenderrrWhatsApp.link('Ravi', t)", text)
    assert unquote(url.split('text=', 1)[1]) == text and '&' not in url.split('text=', 1)[1]
    assert len(url) < 8000, 'Link is too long for WhatsApp: %d' % len(url)
    d.close()


@test
def whatsapp_button_fits_on_all_devices(b, srv):
    for size in SIZES:
        d = Device(b, srv, kind='desktop' if size == 'desktop' else 'android', size=size, clock=OCT_EVENING).open()
        set_split_state(d, FRIENDS, SPLITS, PAYMENTS)
        open_summary(d, 'Ravi')
        r = d.ev("""(() => { const g = id => document.getElementById(id).getBoundingClientRect();
            return { btn: g('waSend'), box: g('sumText'), copy: g('copySum'), w: innerWidth }; })()""")
        assert r['btn']['width'] > 0 and r['btn']['right'] <= r['w'] and r['btn']['left'] >= 0, size + ': button off screen'
        assert abs(r['btn']['width'] - r['box']['width']) < 2, size + ': button should be full width'
        assert r['btn']['top'] >= r['box']['bottom'] and r['btn']['bottom'] <= r['copy']['top'], size + ': button should sit above Copy'
        d.close()


# ---- Money out (debit) and money in (credit) --------------------------------
def summary_card(d):
    return {k: d.page.inner_text('#spx' + k) for k in ('In', 'Out', 'Bal')}


def add_credit(d, amount, cat='Salary', note='', via_switch=False):
    if via_switch:
        d.page.click('#fab')
        d.page.wait_for_selector('#spxSideIn')
        d.page.fill('#fAmt', str(amount))
        d.page.click('#spxSideIn')
    else:
        d.page.click('#spxAddIn')
    d.page.wait_for_selector('#spxAmt')
    if not via_switch:
        d.page.fill('#spxAmt', str(amount))
    d.page.click('#spxCreditCats button.chip:has-text("%s")' % cat)
    if note:
        d.page.fill('#spxNote', note)
    d.page.click('#spxSave')
    d.page.wait_for_timeout(300)


def credits_of(d):
    return json.loads(d.ev("localStorage.getItem('spenderrr.credits') || '{\"e\":[]}'"))['e']


def month_total(d):
    d.page.click('button[data-view=month]')
    d.page.wait_for_timeout(300)
    t = d.ev("document.querySelector('main').innerText")
    return re.search(r'MONTH TOTAL\n(\$[\d,]+\.\d\d)', t).group(1), t


@test
def money_out_has_emi_insurance_and_maid(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    d.page.click('#fab')
    d.page.wait_for_selector('#fAmt')
    chips = d.ev("[...document.querySelectorAll('#sheetRoot button.chip')].map(b => b.textContent.trim())")
    assert chips[-7:] == ['EMI', 'Insurance', 'Maid', 'Fuel', 'LPG', 'Gift', 'Other'] and '+ New' not in chips, chips
    assert d.ev("document.getElementById('spxSideOut').getAttribute('aria-pressed')") == 'true', 'Add form starts on Money out'
    d.page.click('#closeSheet')
    d.add_expense(450, 'EMI')
    d.add_expense(120, 'Insurance')
    d.add_expense(80, 'Maid')
    saved = sorted(e['cat'] for e in json.loads(d.ev("localStorage.getItem('spendlog.v1')"))['e'])
    assert saved == ['emi', 'insurance', 'maid'], saved
    colors = {c: d.ev("(c) => getComputedStyle(document.documentElement).getPropertyValue('--c-' + c).trim()", c) for c in ('emi', 'insurance', 'maid')}
    assert colors == {'emi': '#3F51B5', 'insurance': '#8E44AD', 'maid': '#5D8C1F'}, colors
    total, text = month_total(d)
    assert total == '$650.00' and 'EMI' in text and 'Insurance' in text and 'Maid' in text, text[:400]
    assert summary_card(d)['Out'] == '$650.00'
    d.backup_now()
    wb = open_xlsx(d.folder()[MONTH_XLSX])
    bands = {g[0]: g[1] for g in read_category_sheet(wb['By category'])}
    assert bands == {'EMI': '3F51B5', 'Insurance': '8E44AD', 'Maid': '5D8C1F'}, bands
    d.close()


@test
def money_out_has_fuel_lpg_and_gift(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    d.add_expense(60, 'Fuel')
    d.add_expense(25, 'LPG')
    d.add_expense(40, 'Gift')
    saved = sorted(e['cat'] for e in json.loads(d.ev("localStorage.getItem('spendlog.v1')"))['e'])
    assert saved == ['fuel', 'gift', 'lpg'], saved
    colors = {c: d.ev("(c) => getComputedStyle(document.documentElement).getPropertyValue('--c-' + c).trim()", c) for c in ('fuel', 'lpg', 'gift')}
    assert colors == {'fuel': '#9C6B00', 'lpg': '#0073E6', 'gift': '#BA68C8'}, colors
    total, text = month_total(d)
    assert total == '$125.00' and all(w in text for w in ('Fuel', 'LPG', 'Gift')), text[:400]
    # Money-out Gift and money-in Gift are separate: a gift you give is not income
    assert summary_card(d)['Out'] == '$125.00' and summary_card(d)['In'] == '$0.00'
    d.backup_now()
    wb = open_xlsx(d.folder()[MONTH_XLSX])
    bands = {g[0]: g[1] for g in read_category_sheet(wb['By category'])}
    assert bands == {'Fuel': '9C6B00', 'LPG': '0073E6', 'Gift': 'BA68C8'}, bands
    d.close()


@test
def forms_and_footer_look_right_on_every_device(b, srv):
    """Money out matches Money in (boxed Amount), date and time fit their box, one footer with Saberi RS, Privacy and Version."""
    version = re.search(r'name="app-version" content="([^"]+)"', read('index.html')).group(1)
    for kind in ['iphone', 'android', 'desktop']:
        d = Device(b, srv, kind=kind, clock=OCT_EVENING).open()
        # Footer: printed once
        feet = d.ev("[...document.querySelectorAll('footer')].map(f => f.innerText)")
        assert len(feet) == 1, '%s: one footer, got %s' % (kind, feet)
        f = feet[0]
        assert f.startswith('© 2026 Saberi RS · Spenderrr') and f.count('Privacy Policy') == 1 and f.count('Version') == 1 and f.endswith('Version ' + version), (kind, f)
        # Money out form
        d.page.click('#fab')
        d.page.wait_for_selector('#fAmt')
        d.page.wait_for_timeout(200)
        out = d.ev("""() => { const box = document.querySelector('#sheetRoot .amtfield'), inp = document.getElementById('fAmt'), cs = getComputedStyle(box), ci = getComputedStyle(inp);
          const lab = box.previousElementSibling;
          return { label: lab && lab.textContent, border: cs.borderTopWidth, radius: cs.borderTopLeftRadius, height: Math.round(box.getBoundingClientRect().height),
                   size: ci.fontSize, weight: ci.fontWeight, dollar: getComputedStyle(box.querySelector('span')).display }; }""")
        def date_row(d, field, now=None):
            return d.ev("""([f, n]) => { const i = document.getElementById(f), sheet = i.closest('.sheet'), r = i.getBoundingClientRect(), s = sheet.getBoundingClientRect();
              const pad = parseFloat(getComputedStyle(sheet).paddingRight), cs = getComputedStyle(i);
              const o = { right: r.right, left: r.left, sheetRight: s.right - pad + 0.5, sheetLeft: s.left, appearance: cs.appearance || cs.webkitAppearance, minWidth: cs.minWidth };
              if (n) { const b = document.getElementById(n).getBoundingClientRect(); o.nowLeft = b.left; o.nowRight = b.right; o.nowH = b.height; o.h = r.height; }
              return o; }""", [field, now])
        row = date_row(d, 'fWhen', 'fNow')
        assert row['right'] <= row['nowLeft'] and row['nowRight'] <= row['sheetRight'], '%s: date overlaps Now: %s' % (kind, row)
        assert abs(row['nowH'] - row['h']) <= 1, '%s: Now button should match the date box height: %s' % (kind, row)
        assert row['appearance'] == 'none' and row['minWidth'] == '0px', '%s: date field can shrink on iPhone: %s' % (kind, row)
        d.page.click('#spxSideIn')
        d.page.wait_for_selector('#spxAmt')
        d.page.wait_for_timeout(200)
        inn = d.ev("""() => { const inp = document.getElementById('spxAmt'), ci = getComputedStyle(inp), lab = inp.previousElementSibling;
          return { label: lab && lab.textContent, border: ci.borderTopWidth, radius: ci.borderTopLeftRadius, height: Math.round(inp.getBoundingClientRect().height),
                   size: ci.fontSize, weight: ci.fontWeight }; }""")
        assert out['label'] == inn['label'] == 'Amount', (kind, out, inn)
        for k in ['border', 'radius', 'size', 'weight']:
            assert out[k] == inn[k], '%s: Money out %s %s differs from Money in %s' % (kind, k, out[k], inn[k])
        assert abs(out['height'] - inn['height']) <= 2 and out['dollar'] == 'none', (kind, out, inn)
        row = date_row(d, 'spxWhen')
        assert row['left'] >= row['sheetLeft'] and row['right'] <= row['sheetRight'], '%s: money-in date runs off the edge: %s' % (kind, row)
        d.page.click('#spxCreditClose' if d.page.query_selector('#spxCreditClose') else '.spx-scrim .x')
        d.page.wait_for_timeout(300)
        # Edit expense: same amount box and date row
        d.add_expense(13, 'Lunch')
        d.page.click('main .item')
        d.page.wait_for_selector('#fAmt')
        d.page.wait_for_timeout(200)
        assert d.ev("document.querySelector('#sheetRoot .amtfield').previousElementSibling.textContent") == 'Amount', kind
        row = date_row(d, 'fWhen', 'fNow')
        assert row['right'] <= row['nowLeft'] and row['nowRight'] <= row['sheetRight'], '%s: edit form date overlaps Now: %s' % (kind, row)
        d.close()


@test
def your_own_money_out_category(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    d.page.click('#fab')
    d.page.wait_for_selector('#spxAddDebitCat')
    d.page.fill('#fAmt', '900')
    d.page.fill('#sheetRoot input[placeholder="Where or what"]', 'November rent')
    d.page.click('#spxAddDebitCat')
    d.page.fill('#spxCatName', 'EMI')                        # already exists
    d.page.click('#spxCatSave')
    assert 'already exists' in d.page.inner_text('#spxCatErr')
    d.page.fill('#spxCatName', 'Rent')
    d.page.click('#spxSwatches button[aria-label=Pine]')
    d.page.click('#spxCatSave')
    d.page.wait_for_selector('#sheetRoot button.chip[aria-pressed=true]', timeout=8000)   # app reloads, form reopens
    d.page.wait_for_timeout(300)
    assert d.page.input_value('#fAmt') == '900', 'Typed amount should be kept'
    assert d.page.input_value('#sheetRoot input[placeholder="Where or what"]') == 'November rent'
    picked = d.ev("[...document.querySelectorAll('#sheetRoot button.chip[aria-pressed=true]')].map(b => b.textContent.trim())")
    assert picked == ['Rent'], picked
    d.page.click('#sheetRoot button[type=submit]')
    d.page.wait_for_timeout(300)
    cats = json.loads(d.ev("localStorage.getItem('spenderrr.categories')"))
    rent = cats['debit'][0]
    assert rent['label'] == 'Rent' and rent['color'] == '#004D40' and re.fullmatch(r'c_rent_[a-z0-9]{1,4}', rent['id']), rent
    total, text = month_total(d)
    assert total == '$900.00' and 'Rent' in text and 'Other' not in text.split('BY CATEGORY')[1].split('BY WEEK')[0], text[:400]
    assert d.ev("(id) => SpenderrrLedger.deleteCategory('debit', id).error", rent['id']) == 'Used by 1 entry. Change those first.'
    res = d.ev("SpenderrrLedger.addCategory('debit', '<img src=x onerror=alert(1)>Gym', '#800000')")
    assert res['category']['label'] == 'img src=x onerror=alert(1)Gym'[:24], 'Labels must lose HTML characters: %r' % res
    bad = d.ev("SpenderrrLedger.addCategory('debit', 'Bad', 'red;background:url(x)')")
    assert bad.get('error') == 'Pick a colour.', bad
    d.backup_now()
    backup = json.loads(d.folder()[MONTH_FILE])
    assert [g['category'] for g in backup['byCategory']] == ['Rent']
    d.close()


@test
def money_in_with_switch_list_edit_delete(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    set_split_state(d, ['Ravi'], [('dinner', 60, '2026-10-25 20:15', 'Thai', 'Me', {'Me': 20, 'Ravi': 40}),
                                  ('lunch', 10, '2026-10-24 12:00', '', 'Me', {'Me': 10})])
    total, _ = month_total(d)
    assert summary_card(d)['Out'] == total == '$30.00', 'Money out must match the app (your share of splits)'
    d.page.click('button[data-view=recent]')
    add_credit(d, 2500, 'Salary', 'October pay', via_switch=True)   # amount typed on Money out is carried over
    add_credit(d, 12.5, 'Interest')
    assert summary_card(d) == {'In': '$2,512.50', 'Out': '$30.00', 'Bal': '$2,482.50'}, summary_card(d)
    cr = credits_of(d)
    assert sorted((c['cat'], c['amount']) for c in cr) == [('interest', 12.5), ('salary', 2500)]
    assert len(json.loads(d.ev("localStorage.getItem('spendlog.v1')"))['e']) == 2, 'Money in must not be added as an expense'
    assert month_total(d)[0] == '$30.00', "Money in must not change the app's money out total"
    d.page.click('#spxOpenList')
    d.page.wait_for_selector('#spx-list')
    assert 'october 2026' in d.page.inner_text('#spx-list').lower() and '+$2,500.00' in d.page.inner_text('#spx-list')
    salary = [c for c in cr if c['cat'] == 'salary'][0]
    d.page.click('#spx-list .spx-row[data-id="%s"]' % salary['id'])
    d.page.fill('#spxAmt', '2600')
    d.page.click('#spxSave')
    d.page.wait_for_timeout(300)
    assert summary_card(d)['In'] == '$2,612.50'
    d.page.click('#spxOpenList')
    d.page.click('#spx-list .spx-row[data-id="%s"]' % salary['id'])
    d.page.click('#spxDelete')
    d.page.wait_for_timeout(300)
    assert [c['cat'] for c in credits_of(d)] == ['interest'] and summary_card(d)['In'] == '$12.50'
    d.page.click('#spxAddIn')
    d.page.click('#spxSave')
    assert 'amount above 0' in d.page.inner_text('#spxErr')
    d.page.fill('#spxAmt', '50')
    d.page.click('#spxSave')
    assert 'Pick a category' in d.page.inner_text('#spxErr')
    d.page.click('#spxSideOut')                                       # back to money out keeps the amount
    d.page.wait_for_selector('#fAmt')
    d.page.wait_for_timeout(100)
    assert d.page.input_value('#fAmt') == '50'
    d.close()


@test
def summary_follows_the_month_shown(b, srv):
    d = Device(b, srv, clock=OCT_EVENING).open()
    change(d, ('add', entry('s1', 40, '2026-09-15 10:00')), ('add', entry('o1', 25, '2026-10-05 10:00')))
    d.ev("""localStorage.setItem('spenderrr.credits', JSON.stringify({ e: [
        { id: 'c1', amount: 1000, cat: 'salary', note: '', ts: new Date(2026, 8, 30, 9, 0).getTime(), createdAt: 1 },
        { id: 'c2', amount: 300, cat: 'refund', note: '', ts: new Date(2026, 9, 2, 9, 0).getTime(), createdAt: 2 }] }))""")
    d.reload()
    assert summary_card(d) == {'In': '$300.00', 'Out': '$25.00', 'Bal': '$275.00'}, summary_card(d)
    d.page.click('button[data-view=month]')
    d.page.wait_for_timeout(200)
    d.page.click('main button:has-text("‹")')
    d.page.wait_for_timeout(300)
    assert 'september 2026' in d.page.inner_text('#spx-summary').lower()
    assert summary_card(d) == {'In': '$1,000.00', 'Out': '$40.00', 'Bal': '$960.00'}, summary_card(d)
    d.page.click('button[data-view=split]')
    d.page.wait_for_timeout(200)
    assert not d.page.query_selector('#spx-summary'), 'Summary only on Recent and Month'
    d.close()


@test
def your_own_money_in_category(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    d.page.click('#spxAddIn')
    d.page.fill('#spxAmt', '400')
    d.page.click('#spxAddCreditCat')
    d.page.fill('#spxCatName', 'Freelance')
    d.page.click('#spxSwatches button[aria-label=Tangerine]')
    d.page.click('#spxCatSave')
    d.page.wait_for_timeout(200)
    assert d.page.input_value('#spxAmt') == '400', 'No reload needed for money in categories'
    assert d.ev("document.querySelector('#spxCreditCats button[aria-pressed=true]').textContent.trim()") == 'Freelance'
    d.page.click('#spxSave')
    d.page.wait_for_timeout(300)
    assert summary_card(d)['In'] == '$400.00'
    cats = json.loads(d.ev("localStorage.getItem('spenderrr.categories')"))
    assert [c['label'] for c in cats['credit']] == ['Freelance'] and cats['debit'] == []
    d.page.click('#spxOpenList')
    d.page.click('#spxManage')
    assert 'Freelance' in d.page.inner_text('#spx-manage') and 'Used by 1 entry' in d.page.inner_text('#spx-manage')
    d.close()


@test
def money_in_is_backed_up_merged_and_restored(b, srv):
    drive = FakeDrive()
    a = phone(b, srv, drive, [entry('x1', 30)], paused=False)
    add_credit(a, 2000, 'Salary', 'Phone A pay')
    a.ev("SpenderrrLedger.addCategory('credit', 'Bonus', '#C71585')")
    a.page.clock.pause_at(datetime(2026, 10, 26, 19, 30))
    a.backup_now()
    bp = phone(b, srv, drive, [entry('y1', 15)], paused=False)
    add_credit(bp, 75, 'Refund')
    bp.page.clock.pause_at(datetime(2026, 10, 26, 19, 30))
    bp.backup_now()
    body = json.loads(drive.tree()[MONTH_FILE])
    stored = json.loads(body['data']['spenderrr.credits'])['e']
    assert sorted(c['cat'] for c in stored) == ['refund', 'salary'], "Both phones' money in should be in the backup"
    assert [c['label'] for c in json.loads(body['data']['spenderrr.categories'])['credit']] == ['Bonus']
    assert body['totals'] == {'moneyIn': 2075, 'moneyOut': 45, 'balance': 2030}, body['totals']
    assert [g['category'] for g in body['moneyIn']] == ['Salary', 'Refund']
    load_offer(bp)
    assert sorted(c['cat'] for c in credits_of(bp)) == ['refund', 'salary']
    assert summary_card(bp)['In'] == '$2,075.00'
    salary = [c for c in credits_of(a) if c['cat'] == 'salary'][0]
    a.ev("(id) => { const s = JSON.parse(localStorage.getItem('spenderrr.credits')); s.e = s.e.filter(x => x.id !== id); localStorage.setItem('spenderrr.credits', JSON.stringify(s)); }", salary['id'])
    a.backup_now()                                                # deletion is recorded and reaches phone B
    bp.backup_now()
    load_offer(bp)
    assert [c['cat'] for c in credits_of(bp)] == ['refund'], credits_of(bp)
    wb = open_xlsx(drive.tree()[MONTH_XLSX])
    assert wb.sheetnames == ['By category', 'Summary', 'Money in']
    money_in = read_category_sheet(wb['Money in'])
    assert [(g[0], g[2]) for g in money_in] == [('Refund', 75)], money_in
    summary = {wb['Summary'].cell(r, 2).value: wb['Summary'].cell(r, 4).value for r in range(1, wb['Summary'].max_row + 1)}
    assert summary['Money in'] == 75 and summary['Money out (your share)'] == 45 and summary['Balance'] == 30, summary
    path = os.path.join(tempfile.mkdtemp(), 'spend-log-backup.json')
    open(path, 'w', encoding='utf-8').write(drive.tree()[MONTH_FILE])
    fresh = Device(b, srv, kind='firefox', clock=OCT_EVENING).open()
    fresh.ev('SpendLogBackup.restore()')
    with fresh.page.expect_file_chooser() as fc:
        fresh.page.click('#slb-modal button:has-text("Choose a file")')
    fc.value.set_files(path)
    fresh.page.wait_for_timeout(800)
    fresh.reload()
    assert [c['cat'] for c in credits_of(fresh)] == ['refund'] and summary_card(fresh)['In'] == '$75.00'
    a.close(); bp.close(); fresh.close()


# ---- Categories: add-category prompt and one colour per category ---------
def _lab(h):
    import math
    r, g, b_ = [int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)]
    f = lambda c: ((c + 0.055) / 1.055) ** 2.4 if c > 0.04045 else c / 12.92
    r, g, b_ = f(r), f(g), f(b_)
    x = (r * 0.4124 + g * 0.3576 + b_ * 0.1805) / 0.95047
    y = r * 0.2126 + g * 0.7152 + b_ * 0.0722
    z = (r * 0.0193 + g * 0.1192 + b_ * 0.9505) / 1.08883
    t = lambda v: v ** (1 / 3) if v > 0.008856 else 7.787 * v + 16 / 116
    return 116 * t(y) - 16, 500 * (t(x) - t(y)), 200 * (t(y) - t(z))


def colour_distance(a, b):
    """CIE76 Delta E: about 2 is just noticeable, above 15 is clearly a different colour."""
    import math
    return math.dist(_lab(a), _lab(b))


def contrast_with_white(h):
    r, g, b_ = [int(h[i:i + 2], 16) / 255 for i in (1, 3, 5)]
    f = lambda c: ((c + 0.055) / 1.055) ** 2.4 if c > 0.04045 else c / 12.92
    return 1.05 / (0.2126 * f(r) + 0.7152 * f(g) + 0.0722 * f(b_) + 0.05)


def custom_cats(d):
    return json.loads(d.ev("localStorage.getItem('spenderrr.categories') || '{\"debit\":[],\"credit\":[]}'"))


@test
def add_category_message_and_button(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    d.page.click('#fab')
    d.page.wait_for_selector('#spxAddDebitCatRow')
    row = d.page.inner_text('#spxAddDebitCatRow')
    assert 'see your category' in row and '+ Add category' in row, row
    chips = d.ev("[...document.querySelectorAll('#sheetRoot button.chip')].map(b => b.textContent.trim())")
    assert '+ New' not in chips, 'The old "+ New" chip should be gone'
    below = d.ev("""(() => { const all = [...document.querySelectorAll('#sheetRoot .chips')]; const last = all[all.length - 1];
        return last.getBoundingClientRect().bottom <= document.getElementById('spxAddDebitCatRow').getBoundingClientRect().top; })()""")
    assert below, 'The prompt should sit under the category chips'
    d.page.click('#spxSideIn')
    d.page.wait_for_selector('#spxAddCreditCatRow')
    assert '+ Add category' in d.page.inner_text('#spxAddCreditCatRow')
    d.page.click('#spxAddCreditCat')
    d.page.wait_for_selector('#spxCatName')
    assert 'Crossed-out colours are already used' in d.page.inner_text('#spx-cat')
    d.close()


@test
def palette_is_clearly_different_from_built_in_colours(b, srv):
    d = Device(b, srv).open()
    palette = d.ev('SpenderrrLedger.palette()')
    built = d.ev('SpenderrrLedger.builtInColours()')
    html, backup = read('index.html'), read('backup.js')
    for cid, label in [('breakfast', 'Breakfast'), ('lunch', 'Lunch'), ('dinner', 'Dinner'), ('snacks', 'Snacks'), ('grocery', 'Grocery'),
                       ('shopping', 'Shopping'), ('household', 'Household'), ('utilities', 'Utilities'), ('transport', 'Travel'),
                       ('other', 'Other'), ('emi', 'EMI'), ('insurance', 'Insurance'), ('maid', 'Maid'),
                       ('fuel', 'Fuel'), ('lpg', 'LPG'), ('gift', 'Gift')]:
        shown = d.ev("(c) => getComputedStyle(document.documentElement).getPropertyValue('--c-' + c).trim().toUpperCase()", cid)
        listed = [x['color'] for x in built if x['label'] == label][0]          # money-out entry comes first (Gift is on both sides)
        assert shown == listed, '%s: app shows %s, colour list says %s' % (label, shown, listed)
        assert ("%s: '%s'" % (cid, listed[1:])) in backup or ("'%s', '%s'" % (cid, label)) in backup, label + ' colour missing in backup.js'
    for c in ['2E7D32', '00897B', '5C6BC0', 'D81B60', '6D4C41']:
        assert c in backup, 'Money in colour %s should match backup.js' % c
    pal = [p['color'] for p in palette]
    assert len(pal) == len(set(pal)) == 16
    for p in pal:
        nearest = min(built, key=lambda x: colour_distance(p, x['color']))
        assert colour_distance(p, nearest['color']) >= 15, '%s is too close to %s (%s)' % (p, nearest['label'], nearest['color'])
        assert contrast_with_white(p) >= 3, '%s is too light for white text' % p
    for i, a in enumerate(pal):
        for c in pal[i + 1:]:
            assert colour_distance(a, c) >= 15, '%s and %s look too alike' % (a, c)
    pre = re.search(r"var PALETTE = \[([^\]]+)\]", html).group(1)
    assert re.findall(r"#[0-9A-F]{6}", pre) == pal, 'index.html must use the same palette as ledger.js'
    d.close()


@test
def every_category_gets_its_own_colour(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    first = d.ev("SpenderrrLedger.addCategory('debit', 'Rent', '#800000')")
    assert first['category']['color'] == '#800000'
    again = d.ev("SpenderrrLedger.addCategory('credit', 'Bonus', '#800000')")
    assert again.get('error') == 'That colour is already used by Rent. Pick another.', again
    assert d.ev("SpenderrrLedger.addCategory('debit', 'Gym', '#3F51B5')").get('error') == 'Pick a colour.', 'Built-in colours cannot be reused'
    d.page.click('#spxAddIn')
    d.page.click('#spxAddCreditCat')
    d.page.wait_for_selector('#spxSwatches')
    maroon = d.page.locator('#spxSwatches button[data-color="#800000"]')
    assert maroon.is_disabled() and maroon.get_attribute('aria-disabled') == 'true' and 'used by Rent' in maroon.get_attribute('aria-label')
    picked = d.ev("document.querySelector('#spxSwatches button[aria-pressed=true]').dataset.color")
    assert picked == '#CD5C5C', 'The first free colour should be picked: %s' % picked
    d.page.fill('#spxCatName', 'Bonus')
    d.page.click('#spxCatSave')
    d.page.wait_for_timeout(200)
    assert [c['color'] for c in custom_cats(d)['credit']] == ['#CD5C5C']
    pal = [p['color'] for p in d.ev('SpenderrrLedger.palette()')]
    for i, col in enumerate(pal[2:]):
        r = d.ev("([n, c]) => SpenderrrLedger.addCategory('debit', n, c)", ['Cat %d' % i, col])
        assert 'category' in r, r
    full = d.ev("SpenderrrLedger.addCategory('debit', 'One more', null)")
    assert full.get('error') == 'All colours are in use. Change or delete a category first.', full
    cats = custom_cats(d)
    colours = [c['color'] for c in cats['debit'] + cats['credit']]
    assert len(colours) == len(set(colours)) == 16, colours
    d.close()


@test
def change_a_category_colour(b, srv):
    d = Device(b, srv, kind='iphone', clock=OCT_EVENING).open()
    rent = d.ev("SpenderrrLedger.addCategory('debit', 'Rent', '#800000')")['category']
    d.ev("SpenderrrLedger.addCategory('credit', 'Bonus', '#CD5C5C')")
    d.reload()
    assert d.ev("getComputedStyle(document.documentElement).getPropertyValue('--c-%s').trim()" % rent['id']) == '#800000'
    d.page.click('#spxOpenList')
    d.page.click('#spxManage')
    d.page.click('button[aria-label="Change colour of Rent"]')
    d.page.wait_for_selector('#spx-colour')
    assert d.ev("document.querySelector('#spxSwatches button[aria-pressed=true]').dataset.color") == '#800000', 'Its own colour is selected'
    assert d.page.locator('#spxSwatches button[data-color="#800000"]').is_enabled(), 'Its own colour stays available'
    assert d.page.locator('#spxSwatches button[data-color="#CD5C5C"]').is_disabled(), "Bonus's colour cannot be taken"
    d.page.click('#spxSwatches button[aria-label=Navy]')
    d.page.click('#spxColourSave')
    d.page.wait_for_timeout(200)
    assert custom_cats(d)['debit'][0]['color'] == '#000080'
    assert d.ev("getComputedStyle(document.documentElement).getPropertyValue('--c-%s').trim()" % rent['id']) == '#000080', 'App colour updates at once'
    bad = d.ev("(id) => SpenderrrLedger.changeColour('debit', id, '#CD5C5C')", rent['id'])
    assert bad.get('error') == 'That colour is already used by Bonus. Pick another.', bad
    assert d.ev("(id) => SpenderrrLedger.changeColour('debit', id, '#2E7D32')", rent['id']).get('error') == 'Pick a colour.', 'Salary green is built in'
    d.reload()
    assert d.ev("getComputedStyle(document.documentElement).getPropertyValue('--c-%s').trim()" % rent['id']) == '#000080', 'Kept after reload'
    d.close()


@test
def same_colour_from_two_phones_is_repaired_the_same_way(b, srv):
    clash = {'debit': [{'id': 'c_gym_b', 'label': 'Gym', 'color': '#800000'}, {'id': 'c_loan_c', 'label': 'Loan', 'color': '#3F51B5'}],
             'credit': [{'id': 'c_bonus_a', 'label': 'Bonus', 'color': '#800000'}]}
    results = []
    for kind in ('iphone', 'android'):                   # two devices, same synced data
        d = Device(b, srv, kind=kind, clock=OCT_EVENING).open()
        d.ev("(c) => localStorage.setItem('spenderrr.categories', JSON.stringify(c))", clash)
        d.reload()
        cats = custom_cats(d)
        mapping = {c['id']: c['color'] for c in cats['debit'] + cats['credit']}
        results.append(mapping)
        for c in cats['debit']:
            shown = d.ev("(id) => getComputedStyle(document.documentElement).getPropertyValue('--c-' + id).trim()", c['id'])
            assert shown == c['color'], 'App colour (%s) must match the saved colour (%s) for %s' % (shown, c['color'], c['label'])
        d.close()
    assert results[0] == results[1], 'Every device must repair colours the same way: %s' % results
    assert results[0] == {'c_bonus_a': '#800000', 'c_gym_b': '#CD5C5C', 'c_loan_c': '#FF5722'}, results[0]
    drive = FakeDrive()                                  # and through a real Drive sync
    a = phone(b, srv, drive, [entry('x1', 5)])
    a.ev("SpenderrrLedger.addCategory('debit', 'Rent', '#800000')")
    a.backup_now()
    bp = phone(b, srv, drive, [entry('y1', 6)])
    bp.ev("SpenderrrLedger.addCategory('debit', 'Gym', '#800000')")    # B hasn't seen Rent yet
    bp.backup_now()
    load_offer(bp)
    a.backup_now()
    load_offer(a)
    ca, cb = custom_cats(a), custom_cats(bp)
    cols = lambda c: {x['label']: x['color'] for x in c['debit']}
    assert cols(ca) == cols(cb), 'Both phones should agree: %s vs %s' % (cols(ca), cols(cb))
    assert len(set(cols(ca).values())) == 2, 'Rent and Gym must end up with different colours: %s' % cols(ca)
    a.close(); bp.close()


# ---- Security ----------------------------------------------------------------
def inline_script_hashes():
    import hashlib
    html = read('index.html')
    return ["'sha256-%s'" % base64.b64encode(hashlib.sha256(s.encode('utf-8')).digest()).decode()
            for s in re.findall(r'<script(?:\s+id="[^"]*")?>([\s\S]*?)</script>', html)]


@test
def security_rule_only_runs_the_apps_own_scripts(b, srv):
    rules = parse_headers_file(os.path.join(APP_DIR, '_headers'))
    hdr = next(r['headers'] for r in rules if r['pattern'] == '/*')
    script_src = re.search(r"script-src ([^;]+);", hdr['Content-Security-Policy']).group(1)
    assert "'unsafe-inline'" not in script_src and "'unsafe-eval'" not in script_src, script_src
    expected = inline_script_hashes()
    assert len(expected) == 2, 'Expected 2 scripts inside index.html, found %d' % len(expected)
    for h in expected:
        assert h in script_src, ('index.html changed: update the sha256 fingerprints in _headers (%s)' % h)
    assert hdr.get('X-Content-Type-Options') == 'nosniff' and hdr.get('X-Frame-Options') == 'DENY'
    d = Device(b, srv).open()
    d.ev("""() => { const i = document.createElement('img'); i.setAttribute('onerror', 'window.__pwned = 1'); i.src = 'data:,x';
                    document.body.appendChild(i); const s = document.createElement('script'); s.textContent = 'window.__pwned2 = 1';
                    document.body.appendChild(s); }""")
    d.page.wait_for_timeout(300)
    assert not d.ev('window.__pwned') and not d.ev('window.__pwned2'), 'Injected code must not run'
    blocked = d.ev('window.__csp')
    assert any('script-src' in x for x in blocked), blocked
    d.ctx.close()


@test
def security_odd_ids_and_broken_items_in_a_backup(b, srv):
    """A backup edited by hand (or by someone else) cannot break the merge, add code, or sneak in bad values."""
    bad_cat = dict(entry('badcat', 7), cat='"><img src=x onerror="window.__xss=1">')
    long_note = entry('longnote', 3, note='n' * 5000)
    split = dict(entry('split1', 30), split={'paidBy': 'Me', 'mode': 'equal', 'my': 10,
                 'parts': [{'name': 'Me', 'amt': 10}, {'name': 'Ravi', 'amt': 20}, {'name': '__proto__', 'amt': 1}, {'name': 'Bad', 'amt': 'x'}]})
    store = {'e': [entry('hasOwnProperty', 10), entry('__proto__', 20), entry('constructor', 30), entry('toString', 40),
                   dict(entry('nan', 1), amount='abc'), dict(entry('huge', 1), amount=1e12), dict(entry('nots', 1), ts='soon'),
                   bad_cat, long_note, split, 'just text', None, [1, 2]],
             's': [{'id': 's1', 'name': 'Ravi', 'amt': 5, 'dir': 'in', 'ts': 1793000000000},
                   {'id': 's2', 'name': 'Ravi', 'amt': 5, 'dir': 'sideways', 'ts': 1793000000000}],
             'p': [{'name': 'Ravi', 'addedAt': 1}, {'name': '__proto__', 'addedAt': 2}, {'name': '', 'addedAt': 3}],
             'm': [], 'constructor': {'x': 1}}
    body = json.dumps({'app': 'Spenderrr', 'format': 2, 'createdAt': '2026-10-20T12:00:00Z',
                       'data': {'spendlog.v1': json.dumps(store)},
                       'deleted': {'e': {'__proto__': 'not a time', 'hasOwnProperty': 5, 'gone': 1793000000000}, 'toString': 'x'}})
    drive = FakeDrive()
    root = drive.add('spend-expense-backup', 'root', FOLDER)
    month = drive.add('October_2026', root, FOLDER)
    drive.add('spend-log-backup.json', month, 'application/json', body)
    d = phone(b, srv, drive, entries=[entry('mine', 5)])          # merges with the file and saves it back
    saved = json.loads(json.loads(drive.tree()[MONTH_FILE])['data']['spendlog.v1'])
    got = {e['id']: e['amount'] for e in saved['e']}
    assert got == {'mine': 5, 'badcat': 7, 'longnote': 3, 'split1': 30}, 'Merge must finish and drop only broken items: %s' % got
    by_id = {e['id']: e for e in saved['e']}
    assert by_id['badcat']['cat'] == 'other' and len(by_id['longnote']['note']) == 500, by_id['badcat']
    assert [p['name'] for p in by_id['split1']['split']['parts']] == ['Me', 'Ravi'], by_id['split1']['split']
    assert [x['id'] for x in saved['s']] == ['s1'] and [x['name'] for x in saved['p']] == ['Ravi'], (saved['s'], saved['p'])
    assert 'constructor' not in saved, 'Built-in names are never written into the app data'
    load_offer(d)
    assert amounts(d) == got, amounts(d)
    for v in ['recent', 'week', 'month', 'split']:
        d.page.click('button[data-view=%s]' % v)
        d.page.wait_for_timeout(200)
    assert not d.ev('window.__xss') and not d.ev('({}).polluted'), 'Nothing from the file may run as code'
    sizes = d.ev("""() => { const ok = JSON.stringify({ data: { 'spendlog.v1': '{}' }, pad: 'x'.repeat(1000) });
        const giant = JSON.stringify({ data: { 'spendlog.v1': '{}' }, pad: 'x'.repeat(21 * 1024 * 1024) });
        return [!!SpendLogBackup.parseBackup(ok), !!SpendLogBackup.parseBackup(giant)]; }""")
    assert sizes == [True, False], 'Files over 20 MB are refused instead of freezing the app: %s' % sizes
    d.close()


@test
def security_backup_files_are_checked_before_use(b, srv):
    evil = {
        'app': 'Spenderrr', 'createdAt': '2026-10-26T12:00:00Z',
        'data': {
            'spendlog.v1': json.dumps({'e': [entry('ok1', 10, note='<img src=x onerror="window.__xss=1">')], 'm': [], 'p': [], 's': [],
                                       '__proto__': {'polluted': 1}}),
            'spenderrr.credits': json.dumps({'e': [
                {'id': 'good', 'amount': 100, 'cat': 'salary', 'note': 'pay\u0000\u0007', 'ts': 1793000000000},
                {'id': 'neg', 'amount': -5, 'cat': 'salary', 'ts': 1793000000000},
                {'id': 'huge', 'amount': 1e12, 'cat': 'salary', 'ts': 1793000000000},
                {'id': 'nan', 'amount': 'abc', 'cat': 'salary', 'ts': 1793000000000},
                {'id': '<script>', 'amount': 5, 'cat': 'salary', 'ts': 1793000000000}]}),
            'spenderrr.categories': json.dumps({'debit': [
                {'id': 'c_ok_1', 'label': '<b onmouseover=alert(1)>Gym</b>', 'color': '#C62828'},
                {'id': 'c_css', 'label': 'CSS', 'color': 'red;background:url(//evil)'},
                {'id': 'not valid id', 'label': 'X', 'color': '#000000'}], 'credit': 'nope'}),
            'spenderrr.contacts': json.dumps({'Ravi': '+19195551234', '__proto__': '+19195550000', 'Bad': 'javascript:alert(1)'}),
            '__spendlog_backup_meta': json.dumps({'drive': True, 'evil': 1}),
            'something-else': 'not ours'}}
    path = os.path.join(tempfile.mkdtemp(), 'spend-log-backup.json')
    open(path, 'w', encoding='utf-8').write(json.dumps(evil))
    d = Device(b, srv, kind='firefox', clock=OCT_EVENING).open()
    d.ev('SpendLogBackup.restore()')
    with d.page.expect_file_chooser() as fc:
        d.page.click('#slb-modal button:has-text("Choose a file")')
    fc.value.set_files(path)
    d.page.wait_for_timeout(800)
    d.reload()
    keys = sorted(d.ev('Object.keys(localStorage)'))
    assert 'something-else' not in keys and not json.loads(d.ev("localStorage.getItem('__spendlog_backup_meta') || '{}'")).get('evil'), keys
    cr = credits_of(d)
    assert [c['id'] for c in cr] == ['good'] and cr[0]['note'] == 'pay', cr
    cats = json.loads(d.ev("localStorage.getItem('spenderrr.categories')"))
    assert cats == {'debit': [{'id': 'c_ok_1', 'label': 'b onmouseover=alert(1)Gy', 'color': '#C62828'}], 'credit': []}, cats
    assert json.loads(d.ev("localStorage.getItem('spenderrr.contacts')")) == {'Ravi': '+19195551234'}
    assert '__proto__' not in json.loads(d.ev("localStorage.getItem('spendlog.v1')"))
    for v in ['recent', 'week', 'month']:
        d.page.click('button[data-view=%s]' % v)
        d.page.wait_for_timeout(200)
    d.page.click('#spxOpenList')
    d.page.wait_for_timeout(200)
    assert not d.ev('window.__xss') and not d.ev('({}).polluted'), 'Nothing from the file may run as code'
    assert d.ev("getComputedStyle(document.documentElement).getPropertyValue('--c-c_css')") == '', 'Bad colour must be ignored'
    d.close()


@test
def excel_library_is_a_safe_version(b, srv):
    path = os.path.join(APP_DIR, 'xlsx.full.min.js')
    assert os.path.exists(path), 'xlsx.full.min.js is missing: the Excel button would fail with "The Excel tool didn\'t load"'
    text = open(path, encoding='utf-8', errors='ignore').read()
    # The file also contains versions of its parts (e.g. codepage 1.15.0); SheetJS itself is 0.x.y.
    found = [tuple(int(x) for x in m) for m in re.findall(r'version\s*[:=]\s*"(\d+)\.(\d+)\.(\d+)"', text)]
    sheetjs = [f for f in found if f[0] == 0]
    assert sheetjs, 'Could not find the SheetJS version in xlsx.full.min.js'
    v = sheetjs[0]
    assert v >= (0, 20, 2), ('SheetJS %d.%d.%d has known security issues (CVE-2023-30533, CVE-2024-22363). '
                             'Update xlsx.full.min.js to 0.20.2 or newer from cdn.sheetjs.com.' % v)


@test
def every_file_the_app_needs_is_in_the_release(b, srv):
    """Uploading the release folder replaces the whole site, so every file the app asks for must be in it."""
    html, sw, man = read('index.html'), read('sw.js'), json.loads(read('manifest.webmanifest'))
    wanted = set(re.findall(r'(?:src|href)="([^"#:]+?)(?:\?[^"]*)?"', html))
    wanted |= {i['src'] for i in man['icons']}
    wanted |= {f.lstrip('/') for f in re.findall(r"'(/[^']+)'", sw.split('];')[0]) if f != '/'}
    missing = sorted(f for f in wanted if f and not os.path.exists(os.path.join(APP_DIR, f)))
    assert not missing, 'Missing from the release folder: %s' % missing


@test
def excel_report_downloads(b, srv):
    """The Excel button gives one colour-coded Overview sheet in Times New Roman, with money in beside the split column."""
    d = Device(b, srv, kind='android', clock=OCT_EVENING).open()
    set_split_state(d, ['Ravi'], [
        ('lunch', 12.5, '2026-10-20 12:30', 'Sandwich', 'Me', {'Me': 12.5}),
        ('fuel', 40, '2026-10-22 08:00', 'Shell', 'Me', {'Me': 40}),
        ('dinner', 60, '2026-10-25 20:15', 'Thai Palace', 'Me', {'Me': 30, 'Ravi': 30}),
        ('grocery', 99, '2026-09-30 18:00', 'Last month', 'Me', {'Me': 99})])
    d.ev("""() => { const o = JSON.parse(localStorage.getItem('spendlog.v1'));
      o.e.forEach(x => { if (x.split.parts.length === 1) x.split = null; }); localStorage.setItem('spendlog.v1', JSON.stringify(o)); }""")
    d.reload()
    add_credit(d, 1000, 'Salary', 'October pay')
    d.ev('window.showSaveFilePicker = undefined; 0')       # test browser: plain download instead of the save dialog
    d.page.click('#exportBtn')
    d.page.wait_for_selector('#doExport')
    d.page.click('#sheetRoot .opt[data-k=month]')
    with d.page.expect_download(timeout=10000) as dl:
        d.page.click('#doExport')
    assert dl.value.suggested_filename == 'Spend report - This month - 2026-10-26.xlsx', dl.value.suggested_filename
    from openpyxl import load_workbook
    import io; wb = load_workbook(io.BytesIO(open(dl.value.path(), 'rb').read()))
    assert wb.sheetnames == ['Dashboard'], wb.sheetnames
    ws = wb['Dashboard']
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    fonts = {c.font.name for c in cells}
    assert fonts == {'Times New Roman'}, fonts
    text = [str(c.value) for c in cells]
    assert 'October 2026' in text and 'Last month' not in text, 'Only this month: %s' % text[:10]
    at = lambda label: [c for c in cells if c.value == label][0]
    # Cards: money in 1000, money out = your share 12.5 + 40 + 30 = 82.5, balance 917.5, biggest Fuel
    assert ws.cell(at('MONEY IN').row + 1, at('MONEY IN').column).value == 1000
    assert ws.cell(at('MONEY OUT').row + 1, at('MONEY OUT').column).value == 82.5
    assert ws.cell(at('BALANCE').row + 1, at('BALANCE').column).value == 917.5
    assert ws.cell(at('BIGGEST').row + 1, at('BIGGEST').column).value == 'Fuel $40.00'
    assert at('MONEY IN').border.top.color.rgb[-6:] == '2E7D32' and at('MONEY OUT').border.top.color.rgb[-6:] == 'C2412D'
    # Where the money went: biggest first, with coloured bars
    left = [ws.cell(r, 2).value for r in range(at('WHERE THE MONEY WENT').row + 1, at('WHERE THE MONEY WENT').row + 4)]
    assert [str(v).replace('■ ', '') for v in left] == ['Fuel', 'Dinner', 'Lunch'], left
    fuel_bar = ws.cell(at('WHERE THE MONEY WENT').row + 1, 3)
    assert set(fuel_bar.value) == {'█'} and fuel_bar.font.color.rgb[-6:] == '9C6B00', (fuel_bar.value, fuel_bar.font.color.rgb)
    assert str(ws.cell(at('MONEY IN').row, 6).value) != '' and any('Salary' in t for t in text)
    # Entries: Money in sits next to Split, filter buttons on the header, header repeats when printed
    head = at('Split').row
    headers = [c.value for c in ws[head]]
    assert headers[:8] == [None, 'Date and time', 'Category', 'Note', 'Money out', 'Split', 'Money in', 'Balance'], headers
    entries = [[c.value for c in r] for r in ws.iter_rows(min_row=head + 1, max_row=head + 5)]
    assert [e[2] or e[1] for e in entries] == ['Lunch', 'Fuel', 'Dinner', 'Salary', 'Total'], 'Sorted by date and time: %s' % entries
    assert ws.cell(head + 1, 1).fill.fgColor.rgb[-6:] == 'E26C24', 'Colour stripe for Lunch'
    assert ws.cell(head + 2, 3).font.color.rgb[-6:] == '9C6B00', 'Category name in its colour'
    assert entries[3][6] == 1000 and entries[3][4] is None, entries[3]
    assert entries[2][4] == 30 and entries[2][5] == 'Bill $60.00 paid by you · You $30.00, Ravi $30.00', entries[2]
    assert [e[7] for e in entries[:4]] == [-12.5, -52.5, -82.5, 917.5], 'Running balance'
    assert entries[4][4] == 82.5 and entries[4][6] == 1000 and entries[4][7] == 917.5, entries[4]
    assert ws.auto_filter.ref == 'B%d:H%d' % (head, head + 4), ws.auto_filter.ref
    assert ws.print_title_rows.replace('$', '') == '%d:%d' % (head, head), ws.print_title_rows
    assert "didn't load" not in d.page.inner_text('body')
    d.close()
    # iPhone: the report opens the share sheet (Save to Files, WhatsApp, Mail …)
    ip = Device(b, srv, kind='iphone', clock=OCT_EVENING)
    ip.ctx.add_init_script("Object.defineProperty(navigator, 'userAgent', { get: () => 'Mozilla/5.0 (iPhone; CPU iPhone OS 18_0 like Mac OS X)' });")
    ip.open()
    ip.add_expense(5, 'Snacks')
    ip.page.click('#exportBtn')
    ip.page.click('#doExport')
    ip.page.wait_for_timeout(600)
    assert [f['name'] for f in ip.ev('window.__shared')] == ['Spend report - This month - 2026-10-26.xlsx'], ip.ev('window.__shared')
    ip.close()


@test
def excel_report_with_many_entries(b, srv):
    """A long month: every entry is in the sheet (no cap, no scroll box), filter covers them all."""
    d = Device(b, srv, kind='android', clock=OCT_EVENING).open()
    cats = ['breakfast', 'lunch', 'dinner', 'snacks', 'grocery', 'fuel', 'lpg', 'utilities']
    rows = [(cats[i % len(cats)], 1 + i, '2026-10-%02d %02d:%02d' % (1 + i // 5, 8 + i % 5, i % 60), 'Item %d' % (i + 1)) for i in range(120)]
    set_entries(d, rows)
    b64 = d.ev("""async () => { const a = new Uint8Array(await SpendLogBackup.reportBlob('month', 'This month').arrayBuffer());
      let s = ''; for (let i = 0; i < a.length; i += 8192) s += String.fromCharCode.apply(null, a.subarray(i, i + 8192)); return btoa(s); }""")
    import io, base64, openpyxl
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))['Dashboard']
    head = [c for row in ws.iter_rows() for c in row if c.value == 'Split'][0].row
    notes = [ws.cell(r, 4).value for r in range(head + 1, head + 121)]
    assert notes == ['Item %d' % (i + 1) for i in range(120)], 'All 120 entries in time order'
    assert ws.cell(head + 121, 2).value == 'Total' and ws.cell(head + 121, 5).value == sum(range(1, 121)), ws.cell(head + 121, 5).value
    assert ws.auto_filter.ref == 'B%d:H%d' % (head, head + 120), ws.auto_filter.ref
    d.close()


# ---- Updates and offline ---------------------------------------------------
OLD_SW = """const C='spendlog-v4';
self.addEventListener('install',e=>e.waitUntil(caches.open(C).then(c=>c.addAll(['/','/index.html'])).then(()=>self.skipWaiting())));
self.addEventListener('activate',e=>e.waitUntil(self.clients.claim()));
self.addEventListener('fetch',e=>{ if(e.request.method!=='GET')return; e.respondWith(caches.match(e.request,{ignoreSearch:true}).then(r=>r||fetch(e.request))); });"""


@test
def old_saved_copy_updates_and_works_offline(b, srv):
    tmp = tempfile.mkdtemp()
    try:
        for f in os.listdir(APP_DIR):
            if os.path.isfile(os.path.join(APP_DIR, f)):
                shutil.copy(os.path.join(APP_DIR, f), tmp)
        html = read('index.html')
        v = re.search(r'name="app-version" content="([^"]+)"', html).group(1)
        open(os.path.join(tmp, 'index.html'), 'w', encoding='utf-8').write(
            html.replace('content="%s"' % v, 'content="2000.1.1"'))
        open(os.path.join(tmp, 'sw.js'), 'w', encoding='utf-8').write(OLD_SW)
        server = AppServer(tmp)
        d = Device(b, server, sw=True)
        d.page.goto(server.url + '/')
        d.page.wait_for_timeout(2000)
        d.reload()
        ver = "document.querySelector('meta[name=app-version]').content"
        assert d.ev(ver) == '2000.1.1', 'Setup: phone should be stuck on the old copy'
        for f in ['index.html', 'sw.js', 'backup.js']:           # deploy the release
            shutil.copy(os.path.join(APP_DIR, f), tmp)
        for _ in range(3):
            d.reload()
            d.page.wait_for_timeout(1000)
        assert d.ev(ver) == v, 'Phone did not update, shows ' + d.ev(ver)
        assert d.ev('caches.keys()') == ['spendlog-' + v], 'Old saved copy should be removed'
        d.ctx.set_offline(True)
        d.reload()
        assert d.ev(ver) == v and d.page.is_visible('#slb-fab'), 'App should open offline'
        d.ctx.set_offline(False)
        d.page.goto(server.url + '/privacy.html')                # reading the privacy page must not replace the app's offline copy
        d.page.wait_for_timeout(800)
        d.ctx.set_offline(True)
        d.page.goto(server.url + '/')
        d.page.wait_for_selector('#slb-fab', timeout=8000)
        assert d.ev(ver) == v, 'Offline app was replaced by another page'
        d.ctx.set_offline(False)
        d.ctx.close()
        server.stop()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


@test
def offline_file_leaves_google_alone(b, srv):
    sw = read('sw.js')
    assert 'url.origin !== self.location.origin' in sw, 'sw.js must not handle Google or other sites'
    assert "fetch(req, { cache: 'no-store' })" in sw, 'sw.js must load the latest version first'


# --------------------------------------------------------------------------
def main():
    chosen = [t for t in TESTS if not FILTERS or any(f in t.__name__ for f in FILTERS)]
    passed, failed = [], []
    srv = AppServer(APP_DIR)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not HEADED)
        print('Spenderrr tests — %d to run against %s\n' % (len(chosen), APP_DIR))
        for t in chosen:
            name = t.__name__.replace('_', ' ')
            start = time.time()
            try:
                t(browser, srv)
                passed.append(name)
                print('  PASS  %s  (%.1fs)' % (name, time.time() - start))
            except Exception as e:
                failed.append((name, e))
                print('  FAIL  %s\n        %s' % (name, str(e).strip().splitlines()[0] if str(e).strip() else repr(e)))
                if '--verbose' in sys.argv:
                    traceback.print_exc()
        browser.close()
    srv.stop()
    print('\n%d passed, %d failed' % (len(passed), len(failed)))
    sys.exit(1 if failed else 0)


if __name__ == '__main__':
    main()
