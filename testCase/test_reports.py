#!/usr/bin/env python3
"""
Spenderrr – tests for the Excel report periods and month tabs (backup.js)
========================================================================

What is covered:
  - Financial year follows the currency: INR April–March, USD / EUR January–December
  - This week, this month, last month, last 3 months, last 6 months, financial year,
    last financial year, custom From / To dates and all entries
  - Periods longer than a month: Dashboard first, then one tab per month that has entries,
    newest month first; months with no entries get no tab
  - Amounts show the right currency sign (₹, $, €)
  - The report also takes entries from the backup folder and from a picked backup file,
    and entries deleted on this device stay deleted
  - The period picker on Android, iPhone, a small phone and desktop

Uses the same browser set-up as run_tests.py. The clock is fixed so periods never move.

Run:
    python testCase/test_reports.py                 # all tests
    python testCase/test_reports.py custom          # only tests whose name contains "custom"
    python testCase/test_reports.py --headed        # watch the browser
"""
import base64
import io
import json
import os
import sys
import time
import traceback
from datetime import datetime

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
_args = sys.argv[1:]
sys.argv = [sys.argv[0]]                 # run_tests reads sys.argv when imported
import run_tests as T                    # noqa: E402
sys.argv += _args

import openpyxl                          # noqa: E402
from playwright.sync_api import sync_playwright   # noqa: E402

HEADED = '--headed' in sys.argv
FILTERS = [a for a in _args if not a.startswith('--')]
SAT_OCT_31 = datetime(2026, 10, 31, 19, 0)
TESTS = []


def test(fn):
    TESTS.append(fn)
    return fn


# ---- helpers ---------------------------------------------------------------
def device(b, srv, currency='INR', kind='android', clock=SAT_OCT_31):
    d = T.Device(b, srv, kind=kind, clock=clock)
    if kind == 'iphone':
        d.ctx.add_init_script("Object.defineProperty(navigator, 'userAgent', { get: () => 'Mozilla/5.0 (iPhone; CPU iPhone OS 18_0 like Mac OS X)' });")
    d.open()
    if currency:
        d.ev("(c) => localStorage.setItem('spendlog.currency', c)", currency)
        d.reload()
    return d


def set_credits(d, rows):
    """rows: (category id, amount, 'YYYY-MM-DD HH:MM', note)"""
    d.ev("""(rows) => {
      const e = rows.map((r, i) => { const [day, hm] = r[2].split(' '); const [y, mo, dd] = day.split('-').map(Number);
        const [h, mi] = hm.split(':').map(Number); const ts = new Date(y, mo - 1, dd, h, mi).getTime();
        return { id: 'c' + i, amount: r[1], cat: r[0], note: r[3] || '', ts: ts, createdAt: ts }; });
      localStorage.setItem('spenderrr.credits', JSON.stringify({ e: e }));
    }""", [list(r) for r in rows])
    d.reload()


def workbook(d, k, opts=None, from_backup=False):
    """Make the report in the page and open it with openpyxl."""
    b64 = d.ev("""async (a) => {
      const data = a.fromBackup ? await SpendLogBackup.reportData() : undefined;
      const blob = SpendLogBackup.reportBlob(a.k, undefined, undefined, a.opts, data);
      const u = new Uint8Array(await blob.arrayBuffer()); let s = '';
      for (let i = 0; i < u.length; i += 8192) s += String.fromCharCode.apply(null, u.subarray(i, i + 8192));
      return btoa(s); }""", {'k': k, 'opts': opts, 'fromBackup': from_backup})
    return openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))


def cells(ws):
    return [c for row in ws.iter_rows() for c in row if c.value is not None]


def find(ws, label):
    return [c for c in cells(ws) if c.value == label][0]


def card(ws, label):
    c = find(ws, label)
    return ws.cell(c.row + 1, c.column).value


def notes(ws):
    """Notes column of the entries table, top to bottom (Total row excluded)."""
    head = find(ws, 'Split').row
    out = []
    for r in range(head + 1, ws.max_row + 1):
        if ws.cell(r, 2).value == 'Total':
            break
        if ws.cell(r, 3).value is not None:
            out.append(ws.cell(r, 4).value)
    return out


def period(d, k, opts=None):
    return d.ev("""(a) => { const p = SpendLogBackup.reportPeriod(a.k, new Date(), a.opts); if (!p) return null;
      const f = t => { const x = new Date(t); return x.getFullYear() + '-' + String(x.getMonth() + 1).padStart(2, '0') + '-' + String(x.getDate()).padStart(2, '0'); };
      return [f(p.start), f(p.end - 1)]; }""", {'k': k, 'opts': opts})


def open_sheet(d):
    d.page.click('#exportBtn')
    d.page.wait_for_selector('#sheetRoot .slr-opts')


def picker(d):
    return d.ev("""[...document.querySelectorAll('#sheetRoot .slr-opts .opt')].map(b => ({
      k: b.dataset.k, label: b.querySelector('span').textContent.replace(/\\u2011/g, '-'),
      text: b.innerText.replace(/\\u2011/g, '-'), on: b.getAttribute('aria-pressed') === 'true' }))""")


def download(d):
    d.ev('window.showSaveFilePicker = undefined; 0')
    with d.page.expect_download(timeout=10000) as dl:
        d.page.click('#doExport')
    return dl.value.suggested_filename, openpyxl.load_workbook(io.BytesIO(open(dl.value.path(), 'rb').read()))


YEAR_OF_ENTRIES = [
    ('lunch', 120, '2025-04-03 13:00', 'Apr 25 lunch'),
    ('grocery', 2500, '2026-01-10 18:00', 'Jan big shop'),
    ('snacks', 40, '2026-01-20 16:00', 'Jan tea'),
    ('fuel', 1800, '2026-03-05 09:00', 'Mar fuel'),
    ('dinner', 650, '2026-03-20 20:30', 'Mar dinner'),
    ('utilities', 999, '2026-08-12 10:00', 'Aug power'),
    ('lunch', 180, '2026-10-02 13:00', 'Oct lunch'),
    ('shopping', 3200, '2026-10-30 17:00', 'Oct shirt'),
]


# ---- periods ---------------------------------------------------------------
@test
def financial_year_follows_the_currency(b, srv):
    d = device(b, srv, 'INR')
    assert period(d, 'fy') == ['2026-04-01', '2027-03-31'], period(d, 'fy')
    assert period(d, 'lastfy') == ['2025-04-01', '2026-03-31'], period(d, 'lastfy')
    d.page.clock.set_system_time(datetime(2027, 2, 15, 10, 0))     # still FY 2026-27 in February
    assert period(d, 'fy') == ['2026-04-01', '2027-03-31'], period(d, 'fy')
    d.page.clock.set_system_time(datetime(2027, 4, 1, 10, 0))      # 1 April starts a new one
    assert period(d, 'fy') == ['2027-04-01', '2028-03-31'], period(d, 'fy')
    assert d.ev("SpendLogBackup.periodLabel('fy', new Date())") == 'Financial year 2027-28'
    for code in ['USD', 'EUR']:
        d.ev("(c) => localStorage.setItem('spendlog.currency', c)", code)
        d.page.clock.set_system_time(SAT_OCT_31)
        assert period(d, 'fy') == ['2026-01-01', '2026-12-31'], (code, period(d, 'fy'))
        assert period(d, 'lastfy') == ['2025-01-01', '2025-12-31'], (code, period(d, 'lastfy'))
        assert d.ev("SpendLogBackup.periodLabel('fy', new Date())") == 'Year 2026'
        assert d.ev("SpendLogBackup.periodLabel('lastfy', new Date())") == 'Last year 2025'
    d.close()


@test
def last_months_week_and_custom_ranges(b, srv):
    d = device(b, srv, 'USD')
    assert period(d, 'week') == ['2026-10-26', '2026-11-01'], period(d, 'week')
    assert period(d, 'month') == ['2026-10-01', '2026-10-31']
    assert period(d, 'lastmonth') == ['2026-09-01', '2026-09-30']
    assert period(d, 'last3') == ['2026-08-01', '2026-10-31'], period(d, 'last3')
    assert period(d, 'last6') == ['2026-05-01', '2026-10-31'], period(d, 'last6')
    d.page.clock.set_system_time(datetime(2026, 2, 10, 10, 0))     # crosses into last year
    assert period(d, 'last3') == ['2025-12-01', '2026-02-28'], period(d, 'last3')
    assert period(d, 'last6') == ['2025-09-01', '2026-02-28'], period(d, 'last6')
    assert period(d, 'custom', {'from': '2026-01-15', 'to': '2026-03-10'}) == ['2026-01-15', '2026-03-10']
    assert period(d, 'custom', {'from': '2026-03-10', 'to': '2026-03-10'}) == ['2026-03-10', '2026-03-10'], 'one day'
    assert period(d, 'custom', {'from': '2026-03-11', 'to': '2026-03-10'}) is None, 'From after To'
    assert period(d, 'custom', {'from': '2026-02-30', 'to': '2026-03-10'}) is None, 'not a real date'
    assert period(d, 'custom', {'from': '', 'to': '2026-03-10'}) is None
    d.close()


# ---- month tabs ------------------------------------------------------------
@test
def month_tabs_newest_first_and_empty_months_skipped(b, srv):
    """Last financial year (Apr 2025 – Mar 2026): entries in April, January and March only."""
    d = device(b, srv, 'INR')
    T.set_entries(d, YEAR_OF_ENTRIES)
    wb = workbook(d, 'lastfy')
    assert wb.sheetnames == ['Dashboard', 'March 2026', 'January 2026', 'April 2025'], wb.sheetnames
    dash = wb['Dashboard']
    assert find(dash, 'Financial year 2025-26') is not None
    assert card(dash, 'MONEY OUT') == 120 + 2500 + 40 + 1800 + 650, card(dash, 'MONEY OUT')
    assert notes(dash) == ['Apr 25 lunch', 'Jan big shop', 'Jan tea', 'Mar fuel', 'Mar dinner'], notes(dash)
    assert notes(wb['March 2026']) == ['Mar fuel', 'Mar dinner'] and card(wb['March 2026'], 'MONEY OUT') == 2450
    assert notes(wb['January 2026']) == ['Jan big shop', 'Jan tea'] and card(wb['January 2026'], 'MONEY OUT') == 2540
    assert notes(wb['April 2025']) == ['Apr 25 lunch']
    assert find(wb['March 2026'], 'March 2026').font.name == 'Times New Roman'
    for name in wb.sheetnames:       # every tab has its own filter buttons over its own entries
        ws = wb[name]
        assert ws.auto_filter.ref and ws.auto_filter.ref.startswith('B%d' % find(ws, 'Split').row), (name, ws.auto_filter.ref)
    d.close()


@test
def current_financial_year_last_3_and_last_6_months(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, YEAR_OF_ENTRIES)
    assert workbook(d, 'fy').sheetnames == ['Dashboard', 'October 2026', 'August 2026']
    assert workbook(d, 'last3').sheetnames == ['Dashboard', 'October 2026', 'August 2026']
    assert workbook(d, 'last6').sheetnames == ['Dashboard', 'October 2026', 'August 2026']
    wb = workbook(d, 'all')
    assert wb.sheetnames == ['Dashboard', 'October 2026', 'August 2026', 'March 2026', 'January 2026', 'April 2025'], wb.sheetnames
    assert card(wb['Dashboard'], 'MONEY OUT') == sum(r[1] for r in YEAR_OF_ENTRIES)
    d.close()


@test
def one_month_periods_keep_a_single_dashboard(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, YEAR_OF_ENTRIES + [('lunch', 75, '2026-09-29 13:00', 'Sep 29'), ('lunch', 80, '2026-10-26 13:00', 'Oct 26')])
    for k in ['week', 'month', 'lastmonth']:
        assert workbook(d, k).sheetnames == ['Dashboard'], k
    assert notes(workbook(d, 'week')['Dashboard']) == ['Oct 26', 'Oct shirt']
    # Custom dates inside one month: no month tab either
    assert workbook(d, 'custom', {'from': '2026-10-01', 'to': '2026-10-15'}).sheetnames == ['Dashboard']
    d.close()


@test
def custom_dates_clip_each_month_tab(b, srv):
    d = device(b, srv, 'USD')
    T.set_entries(d, YEAR_OF_ENTRIES)
    wb = workbook(d, 'custom', {'from': '2026-01-15', 'to': '2026-03-10'})
    assert wb.sheetnames == ['Dashboard', 'March 2026', 'January 2026'], wb.sheetnames
    assert notes(wb['Dashboard']) == ['Jan tea', 'Mar fuel']
    assert notes(wb['January 2026']) == ['Jan tea'], 'Jan 10 is before the From date'
    assert notes(wb['March 2026']) == ['Mar fuel'], 'Mar 20 is after the To date'
    sub = [c.value for c in cells(wb['Dashboard']) if str(c.value).startswith('2026-01-15 to 2026-03-10')]
    assert sub, 'The Dashboard says which dates it covers'
    d.close()


@test
def a_month_with_only_money_in_gets_a_tab(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, [('lunch', 200, '2026-08-05 13:00', 'Aug lunch')])
    set_credits(d, [('salary', 90000, '2026-09-01 09:00', 'Sep pay')])
    wb = workbook(d, 'last3')
    assert wb.sheetnames == ['Dashboard', 'September 2026', 'August 2026'], wb.sheetnames
    assert card(wb['September 2026'], 'MONEY IN') == 90000 and card(wb['September 2026'], 'MONEY OUT') == 0
    assert card(wb['Dashboard'], 'BALANCE') == 89800
    d.close()


@test
def empty_period_and_no_entries_at_all(b, srv):
    d = device(b, srv, 'INR')
    wb = workbook(d, 'fy')
    assert wb.sheetnames == ['Dashboard']
    assert find(wb['Dashboard'], 'No entries in this period')
    T.set_entries(d, YEAR_OF_ENTRIES)
    wb = workbook(d, 'custom', {'from': '2026-05-01', 'to': '2026-07-31'})
    assert wb.sheetnames == ['Dashboard'], 'No entries in May–July: no month tabs'
    d.close()


# ---- currency --------------------------------------------------------------
@test
def amounts_show_the_currency_sign(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, [('grocery', 125000, '2026-10-10 18:00', 'Big shop'), ('lunch', 250, '2026-10-11 13:00', 'Lunch')])
    ws = workbook(d, 'month')['Dashboard']
    assert card(ws, 'BIGGEST') == 'Grocery ₹1,25,000.00', card(ws, 'BIGGEST')
    money_cell = ws.cell(find(ws, 'MONEY OUT').row + 1, find(ws, 'MONEY OUT').column)
    assert '₹' in money_cell.number_format and '$' not in money_cell.number_format, money_cell.number_format
    for code, sign, text in [('USD', '$', 'Grocery $125,000.00'), ('EUR', '€', 'Grocery €125,000.00')]:
        d.ev("(c) => localStorage.setItem('spendlog.currency', c)", code)
        ws = workbook(d, 'month')['Dashboard']
        assert card(ws, 'BIGGEST') == text, (code, card(ws, 'BIGGEST'))
        c = ws.cell(find(ws, 'MONEY OUT').row + 1, find(ws, 'MONEY OUT').column)
        assert sign in c.number_format, (code, c.number_format)
    d.close()


@test
def backup_excel_copy_uses_the_currency_sign(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, [('lunch', 250, '2026-10-11 13:00', 'Lunch')])
    b64 = d.ev("""async () => { const u = new Uint8Array(await SpendLogBackup.excelBlob().arrayBuffer()); let s = '';
      for (let i = 0; i < u.length; i += 8192) s += String.fromCharCode.apply(null, u.subarray(i, i + 8192)); return btoa(s); }""")
    ws = openpyxl.load_workbook(io.BytesIO(base64.b64decode(b64)))['By category']
    fmts = {c.number_format for c in cells(ws) if isinstance(c.value, (int, float)) and c.column == 5}
    assert fmts and all('₹' in f for f in fmts), fmts
    d.close()


# ---- data from backups -----------------------------------------------------
@test
def report_takes_entries_from_the_backup_folder(b, srv):
    """Entries only in the backup folder are in the report; one deleted here since the last backup is not."""
    d = device(b, srv, 'INR')
    T.set_entries(d, [('lunch', 100, '2026-09-10 13:00', 'Kept'), ('dinner', 300, '2026-09-12 20:00', 'Deleted here')])
    d.backup_now()
    files = d.folder()
    backup = json.loads(files[T.MONTH_FILE.replace('October', 'October')])
    store = json.loads(backup['data']['spendlog.v1'])
    other = T.entry('phone2', 450, '2026-08-20 19:00', 'grocery', 'From the other phone')
    store['e'].append(other)
    backup['data']['spendlog.v1'] = json.dumps(store)
    T.put_in_folder(d, T.MONTH_FILE, json.dumps(backup))
    T.change(d, ('del', 't1'))
    wb = workbook(d, 'last3', from_backup=True)
    assert notes(wb['Dashboard']) == ['From the other phone', 'Kept'], notes(wb['Dashboard'])
    assert wb.sheetnames == ['Dashboard', 'September 2026', 'August 2026'], wb.sheetnames
    # Making the report never changes the phone's data or the backup file
    assert set(T.amounts(d)) == {'t0'}, T.amounts(d)
    d.close()


@test
def report_from_the_sheet_uses_a_picked_backup_file(b, srv):
    d = device(b, srv, 'INR', kind='iphone')         # iPhone: no backup folder, so a file is the way
    T.set_entries(d, [('lunch', 100, '2026-10-10 13:00', 'On this phone')])
    path = os.path.join(T.tempfile.mkdtemp(), 'spend-log-backup.json')
    open(path, 'w').write(T.backup_body([T.entry('old1', 700, '2026-06-15 12:00', 'grocery', 'Old phone, June'),
                                         T.entry('old2', 50, '2026-10-01 09:00', 'snacks', 'Old phone, October')]))
    open_sheet(d)
    d.page.set_input_files('#slr-file-in', path)
    d.page.wait_for_selector('#slr-file-note:has-text("spend-log-backup.json")')
    assert '(2 records)' in d.page.inner_text('#slr-file-note')
    d.page.click('#sheetRoot .opt[data-k=last6]')
    d.page.click('#doExport')
    d.page.wait_for_timeout(800)
    shared = d.ev('window.__shared')
    assert [f['name'] for f in shared] == ['Spend report - Last 6 months - 2026-10-31.xlsx'], shared
    assert T.amounts(d) == {'t0': 100}, 'The picked file is only read, never restored'
    d.page.click('#closeSheet')
    d.page.click('#exportBtn')
    d.page.wait_for_selector('#slr-file-pick')
    assert 'Add a backup file' in d.page.inner_text('#slr-file-pick'), 'A new sheet starts without the file'
    d.close()
    # The workbook itself (same data, built directly)
    d = device(b, srv, 'INR')
    T.set_entries(d, [('lunch', 100, '2026-10-10 13:00', 'On this phone')])
    open_sheet(d)
    d.page.set_input_files('#slr-file-in', path)
    d.page.wait_for_selector('#slr-file-note:has-text("spend-log-backup.json")')
    d.page.click('#sheetRoot .opt[data-k=last6]')
    name, wb = download(d)
    assert wb.sheetnames == ['Dashboard', 'October 2026', 'June 2026'], wb.sheetnames
    assert notes(wb['October 2026']) == ['Old phone, October', 'On this phone'], notes(wb['October 2026'])
    d.page.click('#closeSheet')
    d.page.click('#exportBtn')
    d.page.wait_for_selector('#slr-file-pick')
    d.page.set_input_files('#slr-file-in', path)
    d.page.wait_for_selector('#slr-file-pick:has-text("Stop using it")')
    d.page.click('#slr-file-pick')
    assert 'Add a backup file' in d.page.inner_text('#slr-file-pick')
    d.close()


@test
def a_file_that_is_not_a_backup_is_refused(b, srv):
    d = device(b, srv, 'INR')
    path = os.path.join(T.tempfile.mkdtemp(), 'notes.json')
    open(path, 'w').write('{"hello": 1}')
    open_sheet(d)
    d.page.set_input_files('#slr-file-in', path)
    d.page.wait_for_timeout(300)
    assert 'not a Spenderrr backup' in d.page.inner_text('#toast')
    assert 'Add a backup file' in d.page.inner_text('#slr-file-pick')
    d.close()


# ---- the picker in the Excel sheet -----------------------------------------
@test
def picker_lists_every_period_for_inr(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, YEAR_OF_ENTRIES)
    open_sheet(d)
    p = picker(d)
    assert [x['k'] for x in p] == ['week', 'month', 'lastmonth', 'last3', 'last6', 'fy', 'lastfy', 'all', 'custom'], p
    assert [x['label'] for x in p] == ['This week', 'This month', 'Last month', 'Last 3 months', 'Last 6 months',
                                       'Financial year 2026-27', 'Last financial year 2025-26', 'All entries', 'Custom dates'], p
    by = {x['k']: x for x in p}
    assert by['month']['on'] and sum(x['on'] for x in p) == 1, 'This month is picked at first'
    assert 'Apr 2026 – Mar 2027' in by['fy']['text'] and 'Apr 2025 – Mar 2026' in by['lastfy']['text']
    assert 'Aug – Oct 2026' in by['last3']['text'] and 'May – Oct 2026' in by['last6']['text']
    assert '2 entries · ₹3,380.00' in by['month']['text'], by['month']['text']
    assert '5 entries · ₹5,110.00' in by['lastfy']['text'], by['lastfy']['text']
    assert not d.page.is_visible('#slr-from'), 'Dates show only for custom'
    d.page.click('#sheetRoot .opt[data-k=lastfy]')
    assert [x['k'] for x in picker(d) if x['on']] == ['lastfy']
    d.page.click('#sheetRoot .opt[data-k=custom]')
    assert d.page.is_visible('#slr-from') and d.page.is_visible('#slr-to')
    assert d.page.input_value('#slr-from') == '2026-10-01' and d.page.input_value('#slr-to') == '2026-10-31'
    d.close()


@test
def picker_uses_calendar_year_for_usd_and_eur(b, srv):
    for code in ['USD', 'EUR']:
        d = device(b, srv, code)
        open_sheet(d)
        by = {x['k']: x for x in picker(d)}
        assert by['fy']['label'] == 'Year 2026' and 'Jan – Dec 2026' in by['fy']['text'], by['fy']
        assert by['lastfy']['label'] == 'Last year 2025' and 'Jan – Dec 2025' in by['lastfy']['text'], by['lastfy']
        assert ('0 entries · %s0.00' % {'USD': '$', 'EUR': '€'}[code]) in by['month']['text'], by['month']['text']
        d.close()


@test
def download_financial_year_from_the_sheet(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, YEAR_OF_ENTRIES)
    open_sheet(d)
    d.page.click('#sheetRoot .opt[data-k=lastfy]')
    name, wb = download(d)
    assert name == 'Spend report - Last financial year 2025-26 - 2026-10-31.xlsx', name
    assert wb.sheetnames == ['Dashboard', 'March 2026', 'January 2026', 'April 2025'], wb.sheetnames
    assert find(wb['Dashboard'], 'Spenderrr · Last financial year 2025-26')
    assert "didn't load" not in d.page.inner_text('body')
    d.close()


@test
def download_custom_dates_from_the_sheet(b, srv):
    d = device(b, srv, 'INR')
    T.set_entries(d, YEAR_OF_ENTRIES)
    open_sheet(d)
    d.page.click('#sheetRoot .opt[data-k=custom]')
    d.page.fill('#slr-from', '2026-01-01')
    d.page.fill('#slr-to', '2026-03-31')
    assert '4 entries · ₹4,990.00' in d.page.inner_text('#sheetRoot .opt[data-k=custom]'), d.page.inner_text('#sheetRoot .opt[data-k=custom]')
    name, wb = download(d)
    assert name == 'Spend report - 2026-01-01 to 2026-03-31 - 2026-10-31.xlsx', name
    assert wb.sheetnames == ['Dashboard', 'March 2026', 'January 2026'], wb.sheetnames
    assert find(wb['Dashboard'], 'Jan – Mar 2026'), 'Whole months read as a month range'
    d.close()


@test
def custom_dates_must_be_in_order(b, srv):
    d = device(b, srv, 'INR')
    open_sheet(d)
    d.page.click('#sheetRoot .opt[data-k=custom]')
    d.page.fill('#slr-from', '2026-03-31')
    d.page.fill('#slr-to', '2026-01-01')
    downloads = []
    d.page.on('download', lambda x: downloads.append(x))
    d.page.click('#doExport')
    d.page.wait_for_timeout(500)
    assert not downloads, 'No file when From is after To'
    assert 'From date must be on or before the To date' in d.page.inner_text('#toast')
    assert 'Pick dates' in d.page.inner_text('#sheetRoot .opt[data-k=custom]')
    d.close()


@test
def picker_keeps_its_choice_when_the_app_redraws_the_sheet(b, srv):
    d = device(b, srv, 'INR')
    open_sheet(d)
    d.page.click('#sheetRoot .opt[data-k=last6]')
    # The app draws its own 4-choice list again (as it does on its own clicks); ours comes back, still on Last 6 months.
    d.ev("""() => { const o = document.querySelector('#sheetRoot .slr-opts'); const app = document.createElement('div'); app.className = 'opts';
      app.innerHTML = '<button class="opt" data-k="week" aria-pressed="false"><span>This week</span></button><button class="opt" data-k="month" aria-pressed="true"><span>This month</span></button>';
      o.replaceWith(app); }""")
    d.page.wait_for_timeout(100)
    assert [x['k'] for x in picker(d) if x['on']] == ['last6'], picker(d)
    assert d.page.locator('#sheetRoot .opts').count() == 1
    d.close()


@test
def picker_fits_on_every_screen(b, srv):
    for kind, size in [('android', None), ('iphone', None), ('android', {'width': 320, 'height': 640}), ('desktop', None)]:
        d = device(b, srv, 'INR', kind=kind)
        if size:
            d.page.set_viewport_size(size)
        open_sheet(d)
        d.page.click('#sheetRoot .opt[data-k=custom]')
        over = d.ev("""() => { const s = document.querySelector('#sheetRoot .sheet');
          const bad = [...s.querySelectorAll('.slr-opts *')].filter(e => { const r = e.getBoundingClientRect(), o = s.getBoundingClientRect();
            return r.width && (r.right > o.right + 1 || r.left < o.left - 1); }).map(e => e.className || e.tagName);
          return { bad: bad, scroll: s.scrollWidth > s.clientWidth + 1 }; }""")
        assert not over['bad'] and not over['scroll'], (kind, size, over)
        d.page.click('#sheetRoot .opt[data-k=fy]')
        assert [x['k'] for x in picker(d) if x['on']] == ['fy'], kind
        d.close()


@test
def iphone_shares_a_last_3_months_report(b, srv):
    d = device(b, srv, 'INR', kind='iphone')
    T.set_entries(d, YEAR_OF_ENTRIES)
    open_sheet(d)
    d.page.click('#sheetRoot .opt[data-k=last3]')
    d.page.click('#doExport')
    d.page.wait_for_timeout(600)
    assert [f['name'] for f in d.ev('window.__shared')] == ['Spend report - Last 3 months - 2026-10-31.xlsx'], d.ev('window.__shared')
    d.close()


def main():
    chosen = [t for t in TESTS if not FILTERS or any(f in t.__name__ for f in FILTERS)]
    passed, failed = [], []
    srv = T.AppServer(T.APP_DIR)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not HEADED)
        print('Spenderrr report tests — %d to run against %s\n' % (len(chosen), T.APP_DIR))
        for t in chosen:
            name = t.__name__.replace('_', ' ')
            start = time.time()
            try:
                t(browser, srv)
                passed.append(name)
                print('  PASS  %s  (%.1fs)' % (name, time.time() - start))
            except Exception as e:
                failed.append((name, e))
                print('  FAIL  %s\n        %s' % (name, str(e).strip().splitlines()[0] if str(e).strip() else repr(e)))
                if '--verbose' in sys.argv:
                    traceback.print_exc()
        browser.close()
    srv.stop()
    print('\n%d passed, %d failed' % (len(passed), len(failed)))
    sys.exit(1 if failed else 0)


if __name__ == '__main__':
    main()
