#!/usr/bin/env python3
"""
Spenderrr – phone app (Android / iPhone) tests
=============================================

Runs the real app bundle (www/, made by build_app.py) in Chromium with a simulated Capacitor
bridge: the phone's files, share sheet, notifications and app launcher are fakes that record
what the app does. This checks native.js and the app bundle without a phone or emulator.

Setup (once):  pip install playwright openpyxl  &&  python -m playwright install chromium
Run:           python testCase/test_native_app.py          (builds www/ first)
               python testCase/test_native_app.py excel    (only tests whose name contains "excel")
"""
import base64
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import traceback
from datetime import datetime
from functools import partial
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer

from playwright.sync_api import sync_playwright
import openpyxl

APP_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
WWW = os.path.join(APP_DIR, 'www')
HEADED = '--headed' in sys.argv
FILTERS = [a for a in sys.argv[1:] if not a.startswith('--')]
OCT = datetime(2026, 10, 26, 17, 0)
MONTH_FILE = 'spend-expense-backup/October_2026/spend-log-backup.json'
MONTH_XLSX = 'spend-expense-backup/October_2026/spend-log-backup.xlsx'
TESTS = []

# --------------------------------------------------------------------------
# Simulated Capacitor bridge (what the Android app provides to the page)
# --------------------------------------------------------------------------
FAKE_BRIDGE = r"""
(function () {
  var S = sessionStorage;                       // survives reloads, like files on a phone
  function load(k, d) { try { return JSON.parse(S.getItem(k)) || d; } catch (e) { return d; } }
  function save(k, v) { S.setItem(k, JSON.stringify(v)); }
  function fsAll() { return load('__fs', {}); }
  function norm(p) { return String(p || '').replace(/^\/+|\/+$/g, ''); }
  function parent(p) { var i = p.lastIndexOf('/'); return i < 0 ? '' : p.slice(0, i); }
  function b64(s) { return btoa(unescape(encodeURIComponent(s))); }
  function fail(msg) { return Promise.reject(new Error(msg)); }
  function denied(dir) { return dir === 'DOCUMENTS' && S.getItem('__denyDocs') === '1'; }
  function log(name, v) { var l = load(name, []); l.push(v); save(name, l); }
  var Filesystem = {
    requestPermissions: function () { return Promise.resolve({ publicStorage: 'granted' }); },
    stat: function (o) {
      var all = fsAll(), d = all[o.directory] || {}, p = norm(o.path);
      if (p === '') return Promise.resolve({ type: 'directory' });
      return d[p] ? Promise.resolve({ type: d[p].type, mtime: d[p].mtime, size: (d[p].data || '').length }) : fail('File does not exist');
    },
    mkdir: function (o) {
      if (denied(o.directory)) return fail('Permission denied');
      var all = fsAll(), d = all[o.directory] = all[o.directory] || {}, p = norm(o.path);
      if (d[p]) return fail('Directory exists');
      var parts = p.split('/'), cur = '';
      for (var i = 0; i < parts.length; i++) {
        cur = cur ? cur + '/' + parts[i] : parts[i];
        if (!d[cur]) { if (i < parts.length - 1 && !o.recursive) return fail('Parent missing'); d[cur] = { type: 'directory', mtime: Date.now() }; }
      }
      save('__fs', all); return Promise.resolve();
    },
    writeFile: function (o) {
      if (denied(o.directory)) return fail('Permission denied');
      var all = fsAll(), d = all[o.directory] = all[o.directory] || {}, p = norm(o.path), par = parent(p);
      if (par && !d[par]) {
        if (!o.recursive) return fail('Parent missing');
        var parts = par.split('/'), cur = '';
        parts.forEach(function (x) { cur = cur ? cur + '/' + x : x; if (!d[cur]) d[cur] = { type: 'directory', mtime: Date.now() }; });
      }
      d[p] = { type: 'file', data: o.encoding === 'utf8' ? b64(o.data) : o.data, mtime: Date.now() };
      save('__fs', all);
      return Promise.resolve({ uri: 'file:///fake/' + o.directory + '/' + p });
    },
    readFile: function (o) {
      var d = fsAll()[o.directory] || {}, f = d[norm(o.path)];
      if (!f || f.type !== 'file') return fail('File does not exist');
      return Promise.resolve({ data: o.encoding === 'utf8' ? decodeURIComponent(escape(atob(f.data))) : f.data });
    },
    readdir: function (o) {
      var d = fsAll()[o.directory] || {}, p = norm(o.path), out = [];
      Object.keys(d).forEach(function (k) { if (parent(k) === p) out.push({ name: k.slice(p ? p.length + 1 : 0), type: d[k].type }); });
      return Promise.resolve({ files: out });
    },
    deleteFile: function (o) { var all = fsAll(); delete (all[o.directory] || {})[norm(o.path)]; save('__fs', all); return Promise.resolve(); },
    rmdir: function (o) { var all = fsAll(), d = all[o.directory] || {}, p = norm(o.path); Object.keys(d).forEach(function (k) { if (k === p || k.indexOf(p + '/') === 0) delete d[k]; }); save('__fs', all); return Promise.resolve(); },
    getUri: function (o) { return Promise.resolve({ uri: 'file:///fake/' + o.directory + '/' + norm(o.path) }); }
  };
  var Share = { share: function (o) {
    log('__shared', o);
    return S.getItem('__cancelShare') === '1' ? fail('Share canceled') : Promise.resolve({});
  } };
  var LocalNotifications = {
    checkPermissions: function () { return Promise.resolve({ display: S.getItem('__notifPerm') || 'prompt' }); },
    requestPermissions: function () { log('__notifAsked', 1); S.setItem('__notifPerm', 'granted'); return Promise.resolve({ display: 'granted' }); },
    schedule: function (o) { log('__scheduled', o.notifications); return Promise.resolve({ notifications: o.notifications.map(function (n) { return { id: n.id }; }) }); },
    cancel: function (o) { log('__cancelled', o.notifications.map(function (n) { return n.id; })); return Promise.resolve(); }
  };
  var AppLauncher = { openUrl: function (o) { log('__opened', o.url); return Promise.resolve({ completed: true }); } };
  var plugins = { Filesystem: Filesystem, Share: Share, LocalNotifications: LocalNotifications, AppLauncher: AppLauncher };
  window.Capacitor = {
    platform: 'android', isNativePlatform: function () { return true; }, getPlatform: function () { return 'android'; },
    Plugins: {},
    nativePromise: function (plugin, method, opts) {
      var p = plugins[plugin];
      if (!p || !p[method]) return Promise.reject(new Error('Not implemented: ' + plugin + '.' + method));
      return p[method](opts || {});
    }
  };
})();
"""


class Server:
    def __init__(self, root):
        class Quiet(SimpleHTTPRequestHandler):
            extensions_map = dict(SimpleHTTPRequestHandler.extensions_map,
                                  **{'.webmanifest': 'application/manifest+json', '.js': 'text/javascript'})

            def log_message(self, *a):
                pass

        self.httpd = ThreadingHTTPServer(('127.0.0.1', 0), partial(Quiet, directory=root))
        self.url = 'http://127.0.0.1:%d' % self.httpd.server_address[1]
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()

    def stop(self):
        self.httpd.shutdown()


class Phone:
    """The installed app on one phone (fresh install unless noted)."""

    def __init__(self, browser, server, clock=OCT, native=True):
        self.server = server
        self.ctx = browser.new_context(viewport={'width': 412, 'height': 860}, accept_downloads=True)
        if native:
            self.ctx.add_init_script(FAKE_BRIDGE)
        self.page = self.ctx.new_page()
        self.errors, self.dialogs, self.answers = [], [], []
        self.page.on('pageerror', lambda e: self.errors.append(str(e)))
        self.page.on('dialog', self._dialog)
        if clock:
            self.page.clock.install(time=clock)

    def _dialog(self, d):
        self.dialogs.append(d.message)
        ans = self.answers.pop(0) if self.answers else True
        d.accept() if ans else d.dismiss()

    def open(self, path='/'):
        self.page.goto(self.server.url + path)
        self.page.wait_for_selector('#slb-fab')
        self.page.wait_for_timeout(300)
        return self

    def reload(self):
        self.page.reload()
        self.page.wait_for_selector('#slb-fab')
        self.page.wait_for_timeout(300)

    def ev(self, js, arg=None):
        return self.page.evaluate(js, arg) if arg is not None else self.page.evaluate(js)

    def log(self, name):
        return json.loads(self.ev("sessionStorage.getItem('%s') || '[]'" % name))

    def files(self, directory='DOCUMENTS'):
        fs = json.loads(self.ev("sessionStorage.getItem('__fs') || '{}'"))
        return {k: v for k, v in fs.get(directory, {}).items() if v['type'] == 'file'}

    def text(self, path, directory='DOCUMENTS'):
        return base64.b64decode(self.files(directory)[path]['data']).decode('utf-8')

    def seed(self, n):
        self.ev("""(n) => {
          const e = [];
          for (let i = 0; i < n; i++) e.push({ id: 'seed' + i, amount: 10 + i, cat: 'lunch', sub: null, note: '',
            ts: Date.now() - i * 3600000, source: 'manual', createdAt: Date.now(), split: null });
          localStorage.setItem('spendlog.v1', JSON.stringify({ e: e, m: [], p: [], s: [] }));
        }""", n)
        self.reload()

    def backup_now(self):
        self.ev('SpendLogBackup.backupNow()')
        self.page.wait_for_timeout(600)

    def close(self):
        assert not self.errors, 'Page errors: %s' % self.errors
        self.ctx.close()


def records_in(text):
    data = json.loads(text)
    store = json.loads(data['data']['spendlog.v1']) if isinstance(data['data']['spendlog.v1'], str) else data['data']['spendlog.v1']
    return len(store['e'])


def test(fn):
    TESTS.append(fn)
    return fn


def read(path):
    return open(os.path.join(APP_DIR, path), encoding='utf-8').read()


# ---- The app bundle --------------------------------------------------------
@test
def bundle_is_the_website_plus_native_js_only(b, srv):
    site, app = read('index.html'), read('www/index.html')
    assert 'native.js' not in site, 'The website must not load native.js'
    assert app.count('<script src="native.js"></script>') == 1
    assert app.index('<script') == app.index('<script src="native.js">'), 'native.js must be the first script'
    assert 'google-client-id' not in app and 'google-client-id' in site
    expected = re.sub(r'\s*<meta name="google-client-id" content="[^"]*">', '', site.replace('<meta charset="utf-8">', '<meta charset="utf-8"> <script src="native.js"></script>', 1))
    assert app == expected, 'The app page differs from the website by more than native.js and the Drive id'
    for f in ['backup.js', 'ledger.js', 'reminders.js', 'whatsapp.js', 'xlsx.full.min.js', 'privacy.html']:
        assert read('www/' + f) == read(f), f + ' must be copied unchanged'
    for f in ['sw.js', '_headers', 't']:
        assert not os.path.exists(os.path.join(WWW, f)), f + ' does not belong in the app'
    gradle = read('android/app/build.gradle')
    version = re.search(r'name="app-version" content="([^"]+)"', site).group(1)
    assert 'versionName "%s"' % version in gradle, 'Android versionName must match the app version'


@test
def native_js_does_nothing_on_the_website(b, srv):
    p = Phone(b, srv, native=False).open()
    assert p.ev('typeof window.SpenderrrNative') == 'undefined'
    assert not p.ev("document.documentElement.classList.contains('spx-native')")
    p.close()


@test
def app_starts_offline_without_errors(b, srv):
    p = Phone(b, srv)
    p.ctx.route('**/*', lambda r: r.abort() if not r.request.url.startswith(srv.url) else r.continue_())
    p.open()
    assert p.ev("document.documentElement.classList.contains('spx-native')")
    assert p.ev('SpenderrrNative.platform') == 'android'
    assert p.ev("document.querySelector('meta[name=google-client-id]') === null")
    assert p.ev('navigator.serviceWorker.getRegistrations().then(r => r.length)') == 0
    assert p.ev("navigator.serviceWorker.register('/sw.js').then(r => typeof r.update)") == 'function'
    p.page.wait_for_timeout(300)
    assert p.ev('navigator.serviceWorker.getRegistrations().then(r => r.length)') == 0, 'No real service worker in the app'
    p.close()


@test
def google_drive_is_hidden_and_device_folder_is_offered(b, srv):
    p = Phone(b, srv).open()
    p.page.click('#slb-fab')
    p.page.wait_for_timeout(200)
    panel = p.page.inner_text('#slb-panel') if p.page.query_selector('#slb-panel') else p.page.inner_text('body')
    assert 'Google Drive' not in panel, panel
    p.ev('SpendLogBackup.restore()')
    p.page.wait_for_selector('#slb-modal')
    modal = p.page.inner_text('#slb-modal')
    assert 'Google Drive' not in modal and 'Backup folder on this device' in modal, modal
    p.close()


# ---- Folder backups go to the phone's Documents folder ---------------------
@test
def backup_goes_to_documents_month_folder(b, srv):
    p = Phone(b, srv).open()
    p.seed(3)
    p.backup_now()
    files = p.files()
    assert MONTH_FILE in files and MONTH_XLSX in files, sorted(files)
    assert records_in(p.text(MONTH_FILE)) == 3
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(files[MONTH_XLSX]['data'])))
    assert wb.sheetnames, 'Excel copy opens'
    p.close()


@test
def remembered_folder_works_after_reopening_the_app(b, srv):
    p = Phone(b, srv).open()
    p.seed(2)
    p.backup_now()
    p.ev("""() => { const o = JSON.parse(localStorage.getItem('spendlog.v1'));
      o.e.push({ id: 'new1', amount: 5, cat: 'snacks', sub: null, note: '', ts: Date.now(), source: 'manual', createdAt: Date.now(), split: null });
      localStorage.setItem('spendlog.v1', JSON.stringify(o)); }""")
    p.reload()
    p.backup_now()
    assert records_in(p.text(MONTH_FILE)) == 3, 'Same file updated after reopening, using the remembered folder'
    assert len([f for f in p.files() if f.endswith('.json')]) == 1
    p.close()


@test
def restore_on_a_reinstalled_app_reads_the_documents_folder(b, srv):
    p = Phone(b, srv).open()
    p.seed(4)
    p.backup_now()
    # Reinstall: app data gone, files in Documents stay.
    p.ev("""async () => { localStorage.clear();
      await new Promise(r => { const q = indexedDB.deleteDatabase('spendlog-backup'); q.onsuccess = q.onerror = q.onblocked = r; }); }""")
    p.reload()
    assert p.ev('SpendLogBackup.records()') == 0
    p.ev('SpendLogBackup.restore()')
    p.page.click('#slb-modal button:has-text("Backup folder on this device")')
    p.page.wait_for_selector('#slb-modal .slb-item')
    p.page.click('#slb-modal .slb-item >> nth=0')
    p.page.wait_for_timeout(1500)
    try:
        p.page.wait_for_selector('#slb-fab', timeout=5000)
    except Exception:
        pass
    assert p.ev('SpendLogBackup.records()') == 4, p.dialogs
    p.close()


@test
def documents_not_writable_falls_back_to_app_storage(b, srv):
    p = Phone(b, srv)
    p.ctx.add_init_script("sessionStorage.setItem('__denyDocs', '1')")
    p.open()
    p.seed(2)
    p.backup_now()
    assert not p.files('DOCUMENTS')
    assert MONTH_FILE in p.files('DATA'), sorted(p.files('DATA'))
    assert p.ev("localStorage.getItem('spenderrr.native.backupDir')") == 'DATA'
    p.close()


# ---- Excel report, downloads and the share sheet ---------------------------
@test
def excel_report_is_saved_and_shared(b, srv):
    p = Phone(b, srv).open()
    p.seed(3)
    p.page.click('#exportBtn')
    p.page.wait_for_selector('#doExport')
    p.page.click('#sheetRoot .opt[data-k=month]')
    p.page.click('#doExport')
    p.page.wait_for_timeout(1500)
    name = 'Spend report - This month - 2026-10-26.xlsx'
    saved = p.files()
    assert 'Spenderrr/' + name in saved, sorted(saved)
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(saved['Spenderrr/' + name]['data'])))
    assert wb.sheetnames == ['Dashboard'], wb.sheetnames
    shared = p.log('__shared')
    assert len(shared) == 1 and shared[0]['files'][0].endswith('/' + name), shared
    assert shared[0]['files'][0].startswith('file:///fake/CACHE/'), 'Shared from the app cache'
    assert 'Saved to Documents' in p.page.inner_text('#spx-native-toast')
    p.close()


@test
def any_download_becomes_save_and_share(b, srv):
    p = Phone(b, srv).open()
    p.ev("""() => { const a = document.createElement('a');
      a.href = URL.createObjectURL(new Blob(['BEGIN:VCALENDAR\\nEND:VCALENDAR'], { type: 'text/calendar' }));
      a.download = 'Electricity-reminder.ics'; document.body.appendChild(a); a.click(); a.remove(); }""")
    p.page.wait_for_timeout(800)
    assert p.text('Spenderrr/Electricity-reminder.ics').startswith('BEGIN:VCALENDAR')
    assert p.log('__shared')[0]['files'][0].endswith('Electricity-reminder.ics')
    p.close()


@test
def share_sheet_with_text_and_cancel(b, srv):
    p = Phone(b, srv).open()
    assert p.ev("navigator.canShare({ files: [new File(['x'], 'a.json')] })") is True
    p.ev("navigator.share({ title: 'Report', text: 'Spent $40 this week' })")
    p.page.wait_for_timeout(200)
    assert p.log('__shared')[0]['text'] == 'Spent $40 this week'
    p.ev("sessionStorage.setItem('__cancelShare', '1')")
    err = p.ev("navigator.share({ text: 'x' }).then(() => 'ok', e => e.name)")
    assert err == 'AbortError', err
    p.close()


# ---- Other apps -------------------------------------------------------------
@test
def whatsapp_email_and_calendar_links_open_in_their_apps(b, srv):
    p = Phone(b, srv).open()
    p.ev("window.open('https://wa.me/?text=hello', '_blank', 'noopener')")
    p.ev("""() => { const a = document.createElement('a'); a.href = 'https://wa.me/15551234567?text=hi'; a.target = '_blank';
      document.body.appendChild(a); a.click(); a.remove(); }""")
    p.ev("""() => { const a = document.createElement('a'); a.href = 'mailto:?subject=Report'; document.body.appendChild(a); a.click(); a.remove(); }""")
    p.ev("SpenderrrNative.openExternal('https://calendar.google.com/calendar/render?action=TEMPLATE')")
    p.page.wait_for_timeout(300)
    opened = p.log('__opened')
    assert opened == ['https://wa.me/?text=hello', 'https://wa.me/15551234567?text=hi', 'mailto:?subject=Report',
                      'https://calendar.google.com/calendar/render?action=TEMPLATE'], opened
    p.close()


@test
def links_inside_the_app_stay_in_the_app(b, srv):
    p = Phone(b, srv).open()
    p.ev("""() => { const a = document.createElement('a'); a.href = 'privacy.html'; a.id = 'pv'; a.textContent = 'Privacy';
      a.style.cssText = 'position:fixed;top:0;left:0;z-index:2147483647;padding:20px;background:#fff'; document.body.appendChild(a); }""")
    p.page.click('#pv')
    p.page.wait_for_load_state()
    assert p.page.url.endswith('/privacy.html'), p.page.url
    assert p.log('__opened') == []
    p.ctx.close()


# ---- Bill reminders as phone notifications ---------------------------------
# Added on the test day (Mon 26 Oct 2026); due on the 28th of every month, reminder 3 days before.
BILL = {'id': 'b_elec1234', 'type': 'electricity', 'name': 'Electricity', 'cat': 'utilities', 'sub': 'Electricity',
        'amount': 1200, 'freq': 'monthly', 'day': 28, 'month': 0, 'remind': 3, 'paid': {},
        'createdAt': int(datetime(2026, 10, 26, 12, 0).timestamp() * 1000)}


def local_times(notifications):
    return [datetime.fromisoformat(n['schedule']['at'].replace('Z', '+00:00')).astimezone().replace(tzinfo=None).strftime('%Y-%m-%d %H:%M')
            for n in notifications]


def add_bill(p, bill, currency='INR'):
    p.ev("""(a) => { localStorage.setItem('spendlog.currency', a.cur); localStorage.setItem('spenderrr.bills', JSON.stringify({ b: [a.bill] })); }""",
         {'bill': bill, 'cur': currency})


@test
def bills_become_phone_notifications(b, srv):
    p = Phone(b, srv).open()
    add_bill(p, BILL)
    p.reload()
    p.page.wait_for_timeout(1500)
    sched = p.log('__scheduled')
    assert sched, 'Notifications were scheduled'
    last = sched[-1]
    # 28 Oct: its "3 days before" (25 Oct) has passed, so remind on the due day; then 3 days before each.
    assert local_times(last) == ['2026-10-28 09:00', '2026-11-25 09:00', '2026-12-25 09:00'], local_times(last)
    assert last[0]['title'] == '💡 Electricity bill due today', last[0]['title']
    second = last[1]
    assert second['title'] == '💡 Electricity bill due in 3 days', second['title']
    assert '₹1,200.00' in second['body'] and 'Nov 28' in second['body'], second['body']
    assert len(p.log('__notifAsked')) == 1, 'Asked for notification permission once'
    assert len(set(n['id'] for n in last)) == 3
    p.close()


@test
def paying_a_bill_moves_its_next_reminder(b, srv):
    p = Phone(b, srv).open()
    add_bill(p, BILL)
    p.reload()
    p.page.wait_for_timeout(1500)
    first = p.log('__scheduled')[-1]
    bill = dict(BILL, paid={'2026-10-28': int(datetime(2026, 10, 26, 17, 0).timestamp() * 1000)})
    add_bill(p, bill)
    p.page.wait_for_timeout(2500)
    now = p.log('__scheduled')[-1]
    assert [n['extra']['due'] for n in now] == ['2026-11-28', '2026-12-28', '2027-01-28'], [n['extra'] for n in now]
    assert local_times(now) == ['2026-11-25 09:00', '2026-12-25 09:00', '2027-01-25 09:00'], local_times(now)
    assert set(p.log('__cancelled')[-1]) == set(n['id'] for n in first), 'Old reminders cancelled first'
    p.close()


@test
def overdue_bill_still_gets_its_coming_reminders(b, srv):
    p = Phone(b, srv).open()
    add_bill(p, dict(BILL, createdAt=int(datetime(2026, 8, 1).timestamp() * 1000)))   # September unpaid: overdue
    p.reload()
    p.page.wait_for_timeout(1500)
    assert local_times(p.log('__scheduled')[-1]) == ['2026-10-28 09:00', '2026-11-25 09:00', '2026-12-25 09:00']
    p.close()


@test
def no_bills_means_no_notification_prompt(b, srv):
    p = Phone(b, srv).open()
    p.page.wait_for_timeout(1500)
    assert p.log('__notifAsked') == [] and p.log('__scheduled') == []
    p.close()


@test
def adding_an_expense_works_in_the_app(b, srv):
    p = Phone(b, srv).open()
    pg = p.page
    pg.click('#fab')
    pg.wait_for_selector('#fAmt')
    pg.fill('#fAmt', '12.5')
    pg.click('#sheetRoot button:has-text("Lunch")')
    pg.click('#sheetRoot button[type=submit]')
    pg.wait_for_timeout(400)
    p.reload()
    assert p.ev('SpendLogBackup.records()') == 1
    p.close()


# ---- Protected (obfuscated) release build ----------------------------------
@test
def protected_build_obfuscates_the_code(b, srv):
    """`build_app.py --protect` obfuscates the authored scripts; assets and index.html are untouched."""
    out = tempfile.mkdtemp(prefix='spx-protected-')
    try:
        subprocess.run([sys.executable, os.path.join(APP_DIR, 'build_app.py'), '--protect', '--out', out],
                       check=True, capture_output=True, text=True)
        # Each authored script: no readable header comment, carries the obfuscator's _0x names.
        for name in ['native.js', 'backup.js', 'ledger.js', 'reminders.js', 'whatsapp.js']:
            code = open(os.path.join(out, name), encoding='utf-8').read()
            assert not code.lstrip().startswith('/*'), name + ' still has its readable header'
            assert '_0x' in code[:200], name + ' is not obfuscated'
        # Some distinctive readable function names that must be gone.
        assert 'function writeToFolder' not in open(os.path.join(out, 'backup.js'), encoding='utf-8').read()
        assert 'function chooseBackupDir' not in open(os.path.join(out, 'native.js'), encoding='utf-8').read()
        assert 'function dueAround' not in open(os.path.join(out, 'reminders.js'), encoding='utf-8').read()
        # index.html (its main script is already obfuscated upstream) and the xlsx library: copied as-is.
        assert open(os.path.join(out, 'index.html'), encoding='utf-8').read() == open(os.path.join(WWW, 'index.html'), encoding='utf-8').read()
        assert open(os.path.join(out, 'xlsx.full.min.js'), encoding='utf-8').read() == read('xlsx.full.min.js')
    finally:
        shutil.rmtree(out, ignore_errors=True)


@test
def protected_build_still_runs_end_to_end(b, srv):
    """The obfuscated bundle must work: backup to Documents and bill notifications both rely on the
    cross-file globals (SpendLogBackup, SpenderrrReminders, SpenderrrNative) surviving obfuscation."""
    out = tempfile.mkdtemp(prefix='spx-protected-run-')
    try:
        subprocess.run([sys.executable, os.path.join(APP_DIR, 'build_app.py'), '--protect', '--out', out],
                       check=True, capture_output=True, text=True)
        rel_srv = Server(out)
        try:
            p = Phone(b, rel_srv).open()
            p.seed(3)
            p.backup_now()
            assert MONTH_FILE in p.files(), sorted(p.files())
            assert records_in(p.text(MONTH_FILE)) == 3
            add_bill(p, BILL)
            p.reload()
            p.page.wait_for_timeout(1500)
            assert p.log('__scheduled'), 'bill notifications scheduled from obfuscated code'
            assert p.log('__scheduled')[-1][0]['title'].endswith('bill due today')
            p.close()
        finally:
            rel_srv.stop()
    finally:
        shutil.rmtree(out, ignore_errors=True)


# --------------------------------------------------------------------------
def main():
    subprocess.run([sys.executable, os.path.join(APP_DIR, 'build_app.py')], check=True)
    chosen = [t for t in TESTS if not FILTERS or any(f in t.__name__ for f in FILTERS)]
    passed, failed = [], []
    srv = Server(WWW)
    with sync_playwright() as p:
        browser = p.chromium.launch(headless=not HEADED)
        print('Spenderrr phone app tests — %d to run against %s\n' % (len(chosen), WWW))
        for t in chosen:
            name = t.__name__.replace('_', ' ')
            start = time.time()
            try:
                t(browser, srv)
                passed.append(name)
                print('  PASS  %s  (%.1fs)' % (name, time.time() - start))
            except Exception as e:
                failed.append((name, e))
                print('  FAIL  %s\n        %s' % (name, str(e).strip().splitlines()[0] if str(e).strip() else repr(e)))
                if '--verbose' in sys.argv:
                    traceback.print_exc()
        browser.close()
    srv.stop()
    print('\n%d passed, %d failed' % (len(passed), len(failed)))
    sys.exit(1 if failed else 0)


if __name__ == '__main__':
    main()
